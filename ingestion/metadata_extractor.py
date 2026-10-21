# backend/ingestion/metadata_extractor.py

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, BinaryIO, Mapping, Optional
import hashlib
import mimetypes
import os
import re


@dataclass(slots=True)
class MetadataExtractionResult:
    """
    Normalized metadata extracted from an uploaded legal document.

    The extractor intentionally focuses on technical/document metadata.
    It does not infer legal meaning, risk, obligations, or compliance.
    """

    filename: Optional[str] = None
    extension: Optional[str] = None
    mime_type: Optional[str] = None
    size_bytes: Optional[int] = None

    sha256: Optional[str] = None

    title: Optional[str] = None
    author: Optional[str] = None
    subject: Optional[str] = None
    creator: Optional[str] = None
    producer: Optional[str] = None

    page_count: Optional[int] = None
    language: Optional[str] = None

    created_at: Optional[datetime] = None
    modified_at: Optional[datetime] = None

    is_encrypted: Optional[bool] = None
    is_scanned: Optional[bool] = None

    custom: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    extracted_at: datetime = field(
        default_factory=lambda: datetime.now(timezone.utc)
    )

    def to_dict(self) -> dict[str, Any]:
        return {
            "filename": self.filename,
            "extension": self.extension,
            "mime_type": self.mime_type,
            "size_bytes": self.size_bytes,
            "sha256": self.sha256,
            "title": self.title,
            "author": self.author,
            "subject": self.subject,
            "creator": self.creator,
            "producer": self.producer,
            "page_count": self.page_count,
            "language": self.language,
            "created_at": (
                self.created_at.isoformat()
                if self.created_at
                else None
            ),
            "modified_at": (
                self.modified_at.isoformat()
                if self.modified_at
                else None
            ),
            "is_encrypted": self.is_encrypted,
            "is_scanned": self.is_scanned,
            "custom": self.custom,
            "warnings": self.warnings,
            "extracted_at": self.extracted_at.isoformat(),
        }


@dataclass(slots=True)
class MetadataExtractorConfig:
    calculate_hash: bool = True
    hash_algorithm: str = "sha256"
    inspect_pdf: bool = True
    inspect_docx: bool = True
    inspect_filesystem: bool = True
    max_hash_bytes: Optional[int] = None


class MetadataExtractionError(Exception):
    """Raised when metadata extraction fails."""

    def __init__(
        self,
        message: str,
        *,
        code: str = "METADATA_EXTRACTION_ERROR",
        cause: Optional[Exception] = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.cause = cause


class MetadataExtractor:
    """
    Extracts technical metadata from documents.

    Supported deep metadata:
        - PDF
        - DOCX

    Generic files still receive:
        - filename
        - extension
        - MIME type
        - size
        - filesystem timestamps
        - content hash

    Design principles:
        - metadata extraction must be deterministic where possible
        - extraction failures should not destroy the document pipeline
        - legal content is never modified
        - parser-specific metadata is kept under `custom`
    """

    SUPPORTED_DEEP_FORMATS = {".pdf", ".docx"}

    MIME_OVERRIDES = {
        ".pdf": "application/pdf",
        ".docx": (
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        ".doc": "application/msword",
        ".txt": "text/plain",
        ".html": "text/html",
        ".htm": "text/html",
        ".json": "application/json",
        ".xml": "application/xml",
        ".csv": "text/csv",
    }

    def __init__(
        self,
        config: Optional[MetadataExtractorConfig] = None,
    ) -> None:
        self.config = config or MetadataExtractorConfig()

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def extract(
        self,
        source: str | Path | bytes | bytearray | BinaryIO,
        *,
        filename: Optional[str] = None,
        mime_type: Optional[str] = None,
        document_metadata: Optional[Mapping[str, Any]] = None,
    ) -> MetadataExtractionResult:
        """
        Extract metadata from a file, bytes object, or binary stream.
        """

        result = MetadataExtractionResult()

        try:
            resolved_filename = self._resolve_filename(
                source,
                filename,
            )

            result.filename = resolved_filename

            extension = self._get_extension(resolved_filename)
            result.extension = extension

            result.mime_type = (
                mime_type
                or self._detect_mime_type(resolved_filename)
            )

            if isinstance(source, (str, Path)):
                path = Path(source)

                if not path.exists():
                    raise MetadataExtractionError(
                        f"File does not exist: {path}",
                        code="FILE_NOT_FOUND",
                    )

                if not path.is_file():
                    raise MetadataExtractionError(
                        f"Source is not a file: {path}",
                        code="INVALID_FILE",
                    )

                if self.config.inspect_filesystem:
                    self._extract_filesystem_metadata(
                        path,
                        result,
                    )

                with path.open("rb") as file_obj:
                    content = file_obj.read()

            elif isinstance(source, (bytes, bytearray)):
                content = bytes(source)
                result.size_bytes = len(content)

            else:
                content = self._read_stream(source)
                result.size_bytes = len(content)

            if result.size_bytes is None:
                result.size_bytes = len(content)

            if self.config.calculate_hash:
                result.sha256 = self._calculate_hash(content)

            # Deep metadata extraction.
            if extension == ".pdf" and self.config.inspect_pdf:
                self._extract_pdf_metadata(
                    content,
                    result,
                )

            elif extension == ".docx" and self.config.inspect_docx:
                self._extract_docx_metadata(
                    content,
                    result,
                )

            # Caller-provided metadata is preserved without overwriting
            # authoritative parser metadata.
            if document_metadata:
                result.custom["pipeline_metadata"] = dict(
                    document_metadata
                )

            return result

        except MetadataExtractionError:
            raise

        except Exception as exc:
            raise MetadataExtractionError(
                "Unexpected metadata extraction failure",
                code="UNEXPECTED_ERROR",
                cause=exc,
            ) from exc

    def extract_from_bytes(
        self,
        content: bytes,
        *,
        filename: str,
        mime_type: Optional[str] = None,
    ) -> MetadataExtractionResult:
        return self.extract(
            content,
            filename=filename,
            mime_type=mime_type,
        )

    def extract_from_file(
        self,
        path: str | Path,
    ) -> MetadataExtractionResult:
        return self.extract(path)

    # ------------------------------------------------------------------
    # Generic metadata
    # ------------------------------------------------------------------

    def _resolve_filename(
        self,
        source: str | Path | bytes | bytearray | BinaryIO,
        filename: Optional[str],
    ) -> Optional[str]:
        if filename:
            return Path(filename).name

        if isinstance(source, (str, Path)):
            return Path(source).name

        source_name = getattr(source, "name", None)

        if source_name:
            return Path(str(source_name)).name

        return None

    @staticmethod
    def _get_extension(filename: Optional[str]) -> str:
        if not filename:
            return ""

        return Path(filename).suffix.lower()

    def _detect_mime_type(
        self,
        filename: Optional[str],
    ) -> Optional[str]:
        if not filename:
            return None

        extension = self._get_extension(filename)

        if extension in self.MIME_OVERRIDES:
            return self.MIME_OVERRIDES[extension]

        mime_type, _ = mimetypes.guess_type(filename)

        return mime_type

    def _read_stream(self, stream: BinaryIO) -> bytes:
        try:
            current_position = stream.tell()
        except Exception:
            current_position = None

        try:
            content = stream.read()

            if isinstance(content, str):
                content = content.encode("utf-8")

            return bytes(content)

        finally:
            if current_position is not None:
                try:
                    stream.seek(current_position)
                except Exception:
                    pass

    # ------------------------------------------------------------------
    # Filesystem metadata
    # ------------------------------------------------------------------

    def _extract_filesystem_metadata(
        self,
        path: Path,
        result: MetadataExtractionResult,
    ) -> None:
        try:
            stat = path.stat()

            result.size_bytes = stat.st_size

            result.created_at = datetime.fromtimestamp(
                getattr(stat, "st_ctime", stat.st_mtime),
                tz=timezone.utc,
            )

            result.modified_at = datetime.fromtimestamp(
                stat.st_mtime,
                tz=timezone.utc,
            )

        except OSError as exc:
            result.warnings.append(
                f"Filesystem metadata unavailable: {exc}"
            )

    # ------------------------------------------------------------------
    # Hashing
    # ------------------------------------------------------------------

    def _calculate_hash(self, content: bytes) -> str:
        algorithm = self.config.hash_algorithm.lower()

        try:
            hasher = hashlib.new(algorithm)
        except ValueError as exc:
            raise MetadataExtractionError(
                f"Unsupported hash algorithm: {algorithm}",
                code="UNSUPPORTED_HASH_ALGORITHM",
                cause=exc,
            ) from exc

        if self.config.max_hash_bytes is not None:
            content = content[: self.config.max_hash_bytes]

        hasher.update(content)

        return hasher.hexdigest()

    # ------------------------------------------------------------------
    # PDF metadata
    # ------------------------------------------------------------------

    def _extract_pdf_metadata(
        self,
        content: bytes,
        result: MetadataExtractionResult,
    ) -> None:
        try:
            from pypdf import PdfReader
        except ImportError:
            result.warnings.append(
                "pypdf is not installed; PDF metadata was not inspected."
            )
            return

        try:
            from io import BytesIO

            reader = PdfReader(BytesIO(content))

            result.page_count = len(reader.pages)

            result.is_encrypted = bool(reader.is_encrypted)

            if reader.metadata:
                self._apply_pdf_metadata(
                    reader.metadata,
                    result,
                )

            # A PDF with pages but no extractable text is commonly a
            # scanned/image-only document. This is only a heuristic.
            if not reader.is_encrypted:
                result.is_scanned = self._estimate_pdf_scan_status(
                    reader
                )

        except Exception as exc:
            result.warnings.append(
                f"PDF metadata extraction failed: {exc}"
            )

    def _apply_pdf_metadata(
        self,
        metadata: Any,
        result: MetadataExtractionResult,
    ) -> None:
        result.title = self._clean_metadata_value(
            getattr(metadata, "title", None)
            or metadata.get("/Title")
        )

        result.author = self._clean_metadata_value(
            getattr(metadata, "author", None)
            or metadata.get("/Author")
        )

        result.subject = self._clean_metadata_value(
            getattr(metadata, "subject", None)
            or metadata.get("/Subject")
        )

        result.creator = self._clean_metadata_value(
            getattr(metadata, "creator", None)
            or metadata.get("/Creator")
        )

        result.producer = self._clean_metadata_value(
            getattr(metadata, "producer", None)
            or metadata.get("/Producer")
        )

        for key, value in metadata.items():
            key_string = str(key)

            if key_string not in {
                "/Title",
                "/Author",
                "/Subject",
                "/Creator",
                "/Producer",
            }:
                result.custom.setdefault(
                    "pdf",
                    {}
                )[key_string] = self._safe_metadata_value(value)

    def _estimate_pdf_scan_status(self, reader: Any) -> bool:
        if not reader.pages:
            return False

        sample_pages = reader.pages[: min(3, len(reader.pages))]

        total_text = 0

        for page in sample_pages:
            try:
                text = page.extract_text() or ""
                total_text += len(text.strip())
            except Exception:
                continue

        return total_text == 0

    # ------------------------------------------------------------------
    # DOCX metadata
    # ------------------------------------------------------------------

    def _extract_docx_metadata(
        self,
        content: bytes,
        result: MetadataExtractionResult,
    ) -> None:
        try:
            from docx import Document
        except ImportError:
            result.warnings.append(
                "python-docx is not installed; DOCX metadata "
                "was not inspected."
            )
            return

        try:
            from io import BytesIO

            document = Document(BytesIO(content))

            properties = document.core_properties

            result.title = self._clean_metadata_value(
                properties.title
            )

            result.author = self._clean_metadata_value(
                properties.author
            )

            result.subject = self._clean_metadata_value(
                properties.subject
            )

            result.creator = self._clean_metadata_value(
                properties.author
            )

            result.created_at = (
                self._normalize_datetime(properties.created)
                or result.created_at
            )

            result.modified_at = (
                self._normalize_datetime(properties.modified)
                or result.modified_at
            )

            # python-docx does not expose an authoritative page count.
            # Store document statistics separately when available.
            paragraph_count = len(document.paragraphs)
            table_count = len(document.tables)
            section_count = len(document.sections)

            result.custom["docx"] = {
                "paragraph_count": paragraph_count,
                "table_count": table_count,
                "section_count": section_count,
                "revision": self._safe_metadata_value(
                    properties.revision
                ),
                "keywords": self._safe_metadata_value(
                    properties.keywords
                ),
                "category": self._safe_metadata_value(
                    properties.category
                ),
                "comments": self._safe_metadata_value(
                    properties.comments
                ),
                "last_modified_by": self._safe_metadata_value(
                    properties.last_modified_by
                ),
            }

        except Exception as exc:
            result.warnings.append(
                f"DOCX metadata extraction failed: {exc}"
            )

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _clean_metadata_value(
        value: Any,
    ) -> Optional[str]:
        if value is None:
            return None

        value = str(value).strip()

        if not value:
            return None

        # Avoid control characters entering API/database metadata.
        value = re.sub(
            r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]",
            "",
            value,
        )

        return value[:10_000] or None

    @staticmethod
    def _safe_metadata_value(value: Any) -> Any:
        if value is None:
            return None

        if isinstance(value, (str, int, float, bool)):
            return value

        if isinstance(value, datetime):
            return value.isoformat()

        try:
            return str(value)
        except Exception:
            return repr(value)

    @staticmethod
    def _normalize_datetime(
        value: Any,
    ) -> Optional[datetime]:
        if value is None:
            return None

        if isinstance(value, datetime):
            if value.tzinfo is None:
                return value.replace(tzinfo=timezone.utc)

            return value.astimezone(timezone.utc)

        return None


def extract_metadata(
    source: str | Path | bytes | bytearray | BinaryIO,
    *,
    filename: Optional[str] = None,
    mime_type: Optional[str] = None,
) -> MetadataExtractionResult:
    """
    Convenience API used by the ingestion pipeline.
    """

    extractor = MetadataExtractor()

    return extractor.extract(
        source,
        filename=filename,
        mime_type=mime_type,
    )
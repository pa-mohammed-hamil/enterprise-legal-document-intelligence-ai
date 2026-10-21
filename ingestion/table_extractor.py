# backend/ingestion/table_extractor.py

from __future__ import annotations

from dataclasses import dataclass, field
from enum import Enum
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence
import re


class TableExtractionStatus(str, Enum):
    SUCCESS = "success"
    PARTIAL = "partial"
    NO_TABLES = "no_tables"
    FAILED = "failed"
    SKIPPED = "skipped"


class TableSourceType(str, Enum):
    PDF = "pdf"
    DOCX = "docx"
    XLSX = "xlsx"
    HTML = "html"
    UNKNOWN = "unknown"


class TableExtractionMethod(str, Enum):
    PYMUPDF = "pymupdf"
    DOCX = "python_docx"
    OPENPYXL = "openpyxl"
    HTML = "html"
    INJECTED = "injected"


@dataclass(slots=True)
class TableCell:
    row: int
    column: int
    text: str = ""

    row_span: int = 1
    column_span: int = 1

    is_header: bool = False

    metadata: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "row": self.row,
            "column": self.column,
            "text": self.text,
            "row_span": self.row_span,
            "column_span": self.column_span,
            "is_header": self.is_header,
            "metadata": self.metadata,
        }


@dataclass(slots=True)
class ExtractedTable:
    table_index: int

    rows: list[list[str]] = field(default_factory=list)
    cells: list[TableCell] = field(default_factory=list)

    page_number: Optional[int] = None

    row_count: int = 0
    column_count: int = 0

    has_header: bool = False

    title: Optional[str] = None

    source_type: TableSourceType = TableSourceType.UNKNOWN
    extraction_method: Optional[str] = None

    bbox: Optional[
        tuple[float, float, float, float]
    ] = None

    sheet_name: Optional[str] = None

    metadata: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    def __post_init__(self) -> None:
        self.row_count = len(self.rows)

        self.column_count = max(
            (
                len(row)
                for row in self.rows
            ),
            default=0,
        )

    def to_markdown(self) -> str:
        """
        Produce a RAG-friendly Markdown representation.
        """

        if not self.rows:
            return ""

        width = self.column_count

        normalized = [
            row + [""] * (width - len(row))
            for row in self.rows
        ]

        if not normalized:
            return ""

        header = normalized[0]

        output = [
            "| "
            + " | ".join(
                self._escape_markdown(value)
                for value in header
            )
            + " |",
            "| "
            + " | ".join(
                "---"
                for _ in range(width)
            )
            + " |",
        ]

        for row in normalized[1:]:
            output.append(
                "| "
                + " | ".join(
                    self._escape_markdown(value)
                    for value in row
                )
                + " |"
            )

        return "\n".join(output)

    def to_text(self) -> str:
        """
        Produce a simple text representation that retains rows.
        """

        output: list[str] = []

        for row in self.rows:
            output.append(
                " | ".join(row)
            )

        return "\n".join(output)

    def to_dict(self) -> dict[str, Any]:
        return {
            "table_index": self.table_index,
            "page_number": self.page_number,
            "rows": self.rows,
            "cells": [
                cell.to_dict()
                for cell in self.cells
            ],
            "row_count": self.row_count,
            "column_count": self.column_count,
            "has_header": self.has_header,
            "title": self.title,
            "source_type": self.source_type.value,
            "extraction_method": self.extraction_method,
            "bbox": self.bbox,
            "sheet_name": self.sheet_name,
            "metadata": self.metadata,
            "warnings": self.warnings,
        }

    @staticmethod
    def _escape_markdown(
        value: str,
    ) -> str:
        return (
            value
            .replace("\\", "\\\\")
            .replace("|", "\\|")
            .replace("\n", "<br>")
        )


@dataclass(slots=True)
class TableExtractionResult:
    status: TableExtractionStatus

    tables: list[ExtractedTable] = field(
        default_factory=list
    )

    table_count: int = 0

    filename: Optional[str] = None
    source_type: TableSourceType = (
        TableSourceType.UNKNOWN
    )

    processed_pages: int = 0
    failed_pages: int = 0

    warnings: list[str] = field(default_factory=list)

    metadata: dict[str, Any] = field(
        default_factory=dict
    )

    def __post_init__(self) -> None:
        self.table_count = len(self.tables)

    def to_dict(self) -> dict[str, Any]:
        return {
            "status": self.status.value,
            "tables": [
                table.to_dict()
                for table in self.tables
            ],
            "table_count": self.table_count,
            "filename": self.filename,
            "source_type": self.source_type.value,
            "processed_pages": self.processed_pages,
            "failed_pages": self.failed_pages,
            "warnings": self.warnings,
            "metadata": self.metadata,
        }


@dataclass(slots=True)
class TableExtractorConfig:
    extract_pdf: bool = True
    extract_docx: bool = True
    extract_xlsx: bool = True
    extract_html: bool = True

    max_pages: Optional[int] = None

    max_tables: Optional[int] = None

    max_rows_per_table: int = 10_000

    max_columns_per_table: int = 200

    max_document_bytes: int = 100 * 1024 * 1024

    detect_headers: bool = True

    remove_empty_rows: bool = True

    remove_empty_columns: bool = True

    normalize_whitespace: bool = True

    preserve_empty_cells: bool = True

    include_cells: bool = True


class TableExtractionError(Exception):

    def __init__(
        self,
        message: str,
        *,
        code: str = "TABLE_EXTRACTION_ERROR",
        cause: Optional[Exception] = None,
    ) -> None:
        super().__init__(message)

        self.code = code
        self.cause = cause


class TableExtractor:
    """
    Structured table extraction for the legal-document pipeline.

    Supported formats
    -----------------
    PDF:
        PyMuPDF table detection.

    DOCX:
        Native Word tables.

    XLSX:
        Worksheets converted into logical tables.

    HTML:
        Native HTML <table> structures.

    Design goal
    -----------
    Tables are preserved as structured data instead of being flattened
    into arbitrary text.

    This is especially important for contracts containing:

        - pricing schedules
        - SLA matrices
        - renewal schedules
        - payment tables
        - liability caps
        - insurance requirements
        - service levels
        - milestone schedules
        - obligation matrices
        - compliance controls
    """

    EXTENSION_MAP = {
        ".pdf": TableSourceType.PDF,
        ".docx": TableSourceType.DOCX,
        ".xlsx": TableSourceType.XLSX,
        ".xlsm": TableSourceType.XLSX,
        ".html": TableSourceType.HTML,
        ".htm": TableSourceType.HTML,
    }

    def __init__(
        self,
        config: Optional[
            TableExtractorConfig
        ] = None,
        *,
        provider: Any = None,
    ) -> None:
        self.config = (
            config
            or TableExtractorConfig()
        )

        self.provider = provider

    # ==============================================================
    # Public API
    # ==============================================================

    def extract(
        self,
        source: (
            str
            | Path
            | bytes
            | bytearray
        ),
        *,
        filename: Optional[str] = None,
        source_type: Optional[
            TableSourceType | str
        ] = None,
    ) -> TableExtractionResult:

        content = self._load_content(
            source
        )

        self._validate_content(
            content
        )

        resolved_filename = (
            self._resolve_filename(
                source,
                filename,
            )
        )

        resolved_type = (
            self._resolve_source_type(
                resolved_filename,
                source_type,
            )
        )

        if self.provider is not None:
            return self._extract_injected(
                content,
                filename=resolved_filename,
                source_type=resolved_type,
            )

        if resolved_type == TableSourceType.PDF:
            return self.extract_pdf(
                content,
                filename=resolved_filename,
            )

        if resolved_type == TableSourceType.DOCX:
            return self.extract_docx(
                content,
                filename=resolved_filename,
            )

        if resolved_type == TableSourceType.XLSX:
            return self.extract_xlsx(
                content,
                filename=resolved_filename,
            )

        if resolved_type == TableSourceType.HTML:
            return self.extract_html(
                content,
                filename=resolved_filename,
            )

        return TableExtractionResult(
            status=TableExtractionStatus.SKIPPED,
            filename=resolved_filename,
            source_type=resolved_type,
            warnings=[
                "Unsupported document type for "
                "table extraction."
            ],
        )

    # ==============================================================
    # PDF
    # ==============================================================

    def extract_pdf(
        self,
        content: bytes,
        *,
        filename: Optional[str] = None,
    ) -> TableExtractionResult:

        if not self.config.extract_pdf:
            return self._disabled_result(
                TableSourceType.PDF,
                filename,
            )

        try:
            import fitz

        except ImportError as exc:
            raise TableExtractionError(
                "PyMuPDF is required for "
                "PDF table extraction.",
                code="PYMUPDF_NOT_INSTALLED",
                cause=exc,
            ) from exc

        document = None

        try:
            document = fitz.open(
                stream=content,
                filetype="pdf",
            )

            if document.needs_pass:
                raise TableExtractionError(
                    "Encrypted PDF requires "
                    "a password.",
                    code="ENCRYPTED_PDF",
                )

            page_count = (
                document.page_count
            )

            limit = page_count

            if (
                self.config.max_pages
                is not None
            ):
                limit = min(
                    limit,
                    self.config.max_pages,
                )

            tables: list[
                ExtractedTable
            ] = []

            warnings: list[str] = []

            failed_pages = 0

            table_index = 1

            for page_index in range(
                limit
            ):
                page_number = (
                    page_index + 1
                )

                try:
                    page = (
                        document.load_page(
                            page_index
                        )
                    )

                    page_tables = (
                        self._extract_pdf_page_tables(
                            page,
                            page_number,
                            table_index,
                        )
                    )

                    tables.extend(
                        page_tables
                    )

                    table_index += len(
                        page_tables
                    )

                    if self._table_limit_reached(
                        tables
                    ):
                        tables = self._apply_table_limit(
                            tables
                        )
                        break

                except Exception as exc:
                    failed_pages += 1

                    warnings.append(
                        f"Table extraction failed "
                        f"on page {page_number}: "
                        f"{exc}"
                    )

            return self._build_result(
                tables,
                filename=filename,
                source_type=(
                    TableSourceType.PDF
                ),
                processed_pages=limit,
                failed_pages=failed_pages,
                warnings=warnings,
                metadata={
                    "document_page_count":
                        page_count,
                    "extraction_method":
                        TableExtractionMethod
                        .PYMUPDF.value,
                },
            )

        except TableExtractionError:
            raise

        except Exception as exc:
            raise TableExtractionError(
                "PDF table extraction failed.",
                code="PDF_TABLE_ERROR",
                cause=exc,
            ) from exc

        finally:
            if document is not None:
                try:
                    document.close()
                except Exception:
                    pass

    def _extract_pdf_page_tables(
        self,
        page: Any,
        page_number: int,
        start_index: int,
    ) -> list[ExtractedTable]:

        if not hasattr(
            page,
            "find_tables",
        ):
            raise TableExtractionError(
                "Installed PyMuPDF version "
                "does not support find_tables().",
                code=(
                    "PDF_TABLE_API_UNAVAILABLE"
                ),
            )

        finder = page.find_tables()

        detected = getattr(
            finder,
            "tables",
            [],
        )

        output: list[
            ExtractedTable
        ] = []

        for offset, table in enumerate(
            detected
        ):
            try:
                raw_rows = (
                    table.extract()
                )

                rows = (
                    self._normalize_rows(
                        raw_rows
                    )
                )

                if not rows:
                    continue

                bbox = self._bbox(
                    getattr(
                        table,
                        "bbox",
                        None,
                    )
                )

                extracted = (
                    self._create_table(
                        table_index=(
                            start_index
                            + offset
                        ),
                        rows=rows,
                        page_number=(
                            page_number
                        ),
                        source_type=(
                            TableSourceType.PDF
                        ),
                        method=(
                            TableExtractionMethod
                            .PYMUPDF.value
                        ),
                        bbox=bbox,
                        metadata={
                            "pdf_page_index":
                                page_number - 1,
                        },
                    )
                )

                output.append(
                    extracted
                )

            except Exception:
                continue

        return output

    # ==============================================================
    # DOCX
    # ==============================================================

    def extract_docx(
        self,
        content: bytes,
        *,
        filename: Optional[str] = None,
    ) -> TableExtractionResult:

        if not self.config.extract_docx:
            return self._disabled_result(
                TableSourceType.DOCX,
                filename,
            )

        try:
            from docx import Document

        except ImportError as exc:
            raise TableExtractionError(
                "python-docx is required for "
                "DOCX table extraction.",
                code=(
                    "PYTHON_DOCX_NOT_INSTALLED"
                ),
                cause=exc,
            ) from exc

        try:
            document = Document(
                BytesIO(content)
            )

            tables: list[
                ExtractedTable
            ] = []

            for index, table in enumerate(
                document.tables,
                start=1,
            ):
                rows: list[
                    list[str]
                ] = []

                for row in table.rows:
                    values = [
                        self._normalize_cell(
                            cell.text
                        )
                        for cell in row.cells
                    ]

                    rows.append(values)

                rows = (
                    self._normalize_rows(
                        rows
                    )
                )

                if not rows:
                    continue

                extracted = (
                    self._create_table(
                        table_index=index,
                        rows=rows,
                        page_number=None,
                        source_type=(
                            TableSourceType.DOCX
                        ),
                        method=(
                            TableExtractionMethod
                            .DOCX.value
                        ),
                        metadata={
                            "physical_page":
                                "unavailable",
                        },
                    )
                )

                tables.append(
                    extracted
                )

                if self._table_limit_reached(
                    tables
                ):
                    break

            return self._build_result(
                self._apply_table_limit(
                    tables
                ),
                filename=filename,
                source_type=(
                    TableSourceType.DOCX
                ),
                processed_pages=1,
                failed_pages=0,
                warnings=[
                    "DOCX table extraction "
                    "preserves logical table "
                    "order; physical page numbers "
                    "are not reliably available."
                ],
                metadata={
                    "document_table_count":
                        len(document.tables),
                    "extraction_method":
                        TableExtractionMethod
                        .DOCX.value,
                },
            )

        except Exception as exc:
            raise TableExtractionError(
                "DOCX table extraction failed.",
                code="DOCX_TABLE_ERROR",
                cause=exc,
            ) from exc

    # ==============================================================
    # XLSX
    # ==============================================================

    def extract_xlsx(
        self,
        content: bytes,
        *,
        filename: Optional[str] = None,
    ) -> TableExtractionResult:

        if not self.config.extract_xlsx:
            return self._disabled_result(
                TableSourceType.XLSX,
                filename,
            )

        try:
            from openpyxl import (
                load_workbook,
            )

        except ImportError as exc:
            raise TableExtractionError(
                "openpyxl is required for "
                "XLSX table extraction.",
                code=(
                    "OPENPYXL_NOT_INSTALLED"
                ),
                cause=exc,
            ) from exc

        workbook = None

        try:
            workbook = load_workbook(
                BytesIO(content),
                read_only=True,
                data_only=True,
            )

            tables: list[
                ExtractedTable
            ] = []

            table_index = 1

            sheet_names = (
                workbook.sheetnames
            )

            if (
                self.config.max_pages
                is not None
            ):
                sheet_names = sheet_names[
                    : self.config.max_pages
                ]

            for sheet_name in sheet_names:
                sheet = workbook[
                    sheet_name
                ]

                rows: list[
                    list[str]
                ] = []

                for row_index, row in enumerate(
                    sheet.iter_rows(
                        values_only=True
                    ),
                    start=1,
                ):
                    if (
                        row_index
                        > self.config
                        .max_rows_per_table
                    ):
                        break

                    values = [
                        self._normalize_cell(
                            value
                        )
                        for value in row[
                            : self.config
                            .max_columns_per_table
                        ]
                    ]

                    rows.append(values)

                rows = (
                    self._normalize_rows(
                        rows
                    )
                )

                if not rows:
                    continue

                table = self._create_table(
                    table_index=table_index,
                    rows=rows,
                    page_number=None,
                    source_type=(
                        TableSourceType.XLSX
                    ),
                    method=(
                        TableExtractionMethod
                        .OPENPYXL.value
                    ),
                    sheet_name=sheet_name,
                    metadata={
                        "worksheet":
                            sheet_name,
                        "logical_page": True,
                    },
                )

                tables.append(table)

                table_index += 1

                if self._table_limit_reached(
                    tables
                ):
                    break

            return self._build_result(
                self._apply_table_limit(
                    tables
                ),
                filename=filename,
                source_type=(
                    TableSourceType.XLSX
                ),
                processed_pages=len(
                    sheet_names
                ),
                failed_pages=0,
                warnings=[
                    "Each XLSX worksheet is "
                    "represented as a logical "
                    "table."
                ],
                metadata={
                    "worksheet_count":
                        len(
                            workbook.sheetnames
                        ),
                    "extraction_method":
                        TableExtractionMethod
                        .OPENPYXL.value,
                },
            )

        except Exception as exc:
            raise TableExtractionError(
                "XLSX table extraction failed.",
                code="XLSX_TABLE_ERROR",
                cause=exc,
            ) from exc

        finally:
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass

    # ==============================================================
    # HTML
    # ==============================================================

    def extract_html(
        self,
        content: bytes,
        *,
        filename: Optional[str] = None,
    ) -> TableExtractionResult:

        if not self.config.extract_html:
            return self._disabled_result(
                TableSourceType.HTML,
                filename,
            )

        try:
            from bs4 import BeautifulSoup

        except ImportError as exc:
            raise TableExtractionError(
                "BeautifulSoup4 is required "
                "for HTML table extraction.",
                code=(
                    "BEAUTIFULSOUP_NOT_INSTALLED"
                ),
                cause=exc,
            ) from exc

        try:
            html = content.decode(
                "utf-8",
                errors="replace",
            )

            soup = BeautifulSoup(
                html,
                "html.parser",
            )

            detected = soup.find_all(
                "table"
            )

            tables: list[
                ExtractedTable
            ] = []

            for index, table_tag in enumerate(
                detected,
                start=1,
            ):
                rows: list[
                    list[str]
                ] = []

                header_rows: set[int] = set()

                for row_index, tr in enumerate(
                    table_tag.find_all("tr")
                ):
                    cells = tr.find_all(
                        ["th", "td"]
                    )

                    values: list[str] = []

                    has_header = False

                    for cell in cells:
                        if cell.name == "th":
                            has_header = True

                        values.append(
                            self._normalize_cell(
                                cell.get_text(
                                    " ",
                                    strip=True,
                                )
                            )
                        )

                    if values:
                        rows.append(values)

                        if has_header:
                            header_rows.add(
                                len(rows) - 1
                            )

                rows = self._normalize_rows(
                    rows
                )

                if not rows:
                    continue

                extracted = self._create_table(
                    table_index=index,
                    rows=rows,
                    page_number=None,
                    source_type=(
                        TableSourceType.HTML
                    ),
                    method=(
                        TableExtractionMethod
                        .HTML.value
                    ),
                    metadata={
                        "html_header_rows":
                            sorted(
                                header_rows
                            ),
                    },
                )

                if 0 in header_rows:
                    extracted.has_header = True

                tables.append(extracted)

                if self._table_limit_reached(
                    tables
                ):
                    break

            return self._build_result(
                self._apply_table_limit(
                    tables
                ),
                filename=filename,
                source_type=(
                    TableSourceType.HTML
                ),
                processed_pages=1,
                failed_pages=0,
                warnings=[],
                metadata={
                    "html_table_count":
                        len(detected),
                    "extraction_method":
                        TableExtractionMethod
                        .HTML.value,
                },
            )

        except Exception as exc:
            raise TableExtractionError(
                "HTML table extraction failed.",
                code="HTML_TABLE_ERROR",
                cause=exc,
            ) from exc

    # ==============================================================
    # Injected provider
    # ==============================================================

    def _extract_injected(
        self,
        content: bytes,
        *,
        filename: Optional[str],
        source_type: TableSourceType,
    ) -> TableExtractionResult:

        try:
            if callable(self.provider):
                response = self.provider(
                    content,
                    filename=filename,
                    source_type=(
                        source_type.value
                    ),
                )

            elif hasattr(
                self.provider,
                "extract_tables",
            ):
                response = (
                    self.provider
                    .extract_tables(
                        content,
                        filename=filename,
                    )
                )

            elif hasattr(
                self.provider,
                "extract",
            ):
                response = (
                    self.provider.extract(
                        content,
                        filename=filename,
                    )
                )

            else:
                raise TableExtractionError(
                    "Injected table provider "
                    "must be callable or expose "
                    "extract()/extract_tables().",
                    code=(
                        "INVALID_TABLE_PROVIDER"
                    ),
                )

            tables = (
                self._normalize_provider_result(
                    response,
                    source_type,
                )
            )

            return self._build_result(
                tables,
                filename=filename,
                source_type=source_type,
                processed_pages=0,
                failed_pages=0,
                warnings=[],
                metadata={
                    "extraction_method":
                        TableExtractionMethod
                        .INJECTED.value,
                },
            )

        except TableExtractionError:
            raise

        except Exception as exc:
            raise TableExtractionError(
                "Injected table extraction "
                "failed.",
                code=(
                    "INJECTED_TABLE_ERROR"
                ),
                cause=exc,
            ) from exc

    def _normalize_provider_result(
        self,
        response: Any,
        source_type: TableSourceType,
    ) -> list[ExtractedTable]:

        if isinstance(
            response,
            TableExtractionResult,
        ):
            return response.tables

        if isinstance(
            response,
            ExtractedTable,
        ):
            return [response]

        if not isinstance(
            response,
            (list, tuple),
        ):
            raise TableExtractionError(
                "Unsupported table provider "
                "response.",
                code=(
                    "INVALID_PROVIDER_RESPONSE"
                ),
            )

        output: list[
            ExtractedTable
        ] = []

        for index, item in enumerate(
            response,
            start=1,
        ):
            if isinstance(
                item,
                ExtractedTable,
            ):
                output.append(item)
                continue

            if isinstance(item, dict):
                rows = item.get(
                    "rows",
                    [],
                )

                output.append(
                    self._create_table(
                        table_index=item.get(
                            "table_index",
                            index,
                        ),
                        rows=self._normalize_rows(
                            rows
                        ),
                        page_number=item.get(
                            "page_number"
                        ),
                        source_type=(
                            source_type
                        ),
                        method=(
                            TableExtractionMethod
                            .INJECTED.value
                        ),
                        bbox=item.get(
                            "bbox"
                        ),
                        sheet_name=item.get(
                            "sheet_name"
                        ),
                        metadata=item.get(
                            "metadata",
                            {},
                        ),
                    )
                )

                continue

            if isinstance(
                item,
                (list, tuple),
            ):
                output.append(
                    self._create_table(
                        table_index=index,
                        rows=self._normalize_rows(
                            item
                        ),
                        page_number=None,
                        source_type=(
                            source_type
                        ),
                        method=(
                            TableExtractionMethod
                            .INJECTED.value
                        ),
                    )
                )

        return output

    # ==============================================================
    # Table construction
    # ==============================================================

    def _create_table(
        self,
        *,
        table_index: int,
        rows: Sequence[
            Sequence[Any]
        ],
        page_number: Optional[int],
        source_type: TableSourceType,
        method: str,
        bbox: Optional[
            tuple[
                float,
                float,
                float,
                float,
            ]
        ] = None,
        sheet_name: Optional[str] = None,
        metadata: Optional[
            dict[str, Any]
        ] = None,
    ) -> ExtractedTable:

        normalized = self._normalize_rows(
            rows
        )

        if (
            len(normalized)
            > self.config.max_rows_per_table
        ):
            normalized = normalized[
                : self.config
                .max_rows_per_table
            ]

        normalized = [
            row[
                : self.config
                .max_columns_per_table
            ]
            for row in normalized
        ]

        has_header = (
            self._detect_header(
                normalized
            )
            if self.config.detect_headers
            else False
        )

        cells: list[
            TableCell
        ] = []

        if self.config.include_cells:
            for row_index, row in enumerate(
                normalized
            ):
                for column_index, value in (
                    enumerate(row)
                ):
                    cells.append(
                        TableCell(
                            row=row_index,
                            column=column_index,
                            text=value,
                            is_header=(
                                has_header
                                and row_index == 0
                            ),
                        )
                    )

        return ExtractedTable(
            table_index=table_index,
            rows=normalized,
            cells=cells,
            page_number=page_number,
            has_header=has_header,
            source_type=source_type,
            extraction_method=method,
            bbox=bbox,
            sheet_name=sheet_name,
            metadata=metadata or {},
        )

    # ==============================================================
    # Header detection
    # ==============================================================

    def _detect_header(
        self,
        rows: list[list[str]],
    ) -> bool:

        if len(rows) < 2:
            return False

        first = rows[0]
        second = rows[1]

        if not first:
            return False

        first_non_empty = [
            value
            for value in first
            if value.strip()
        ]

        if not first_non_empty:
            return False

        # Strong signal: first row is mostly textual labels while the
        # second row contains numbers/dates/values.
        first_numeric = sum(
            self._looks_numeric(value)
            for value in first_non_empty
        )

        second_non_empty = [
            value
            for value in second
            if value.strip()
        ]

        second_numeric = sum(
            self._looks_numeric(value)
            for value in second_non_empty
        )

        first_ratio = (
            first_numeric
            / len(first_non_empty)
        )

        second_ratio = (
            second_numeric
            / len(second_non_empty)
            if second_non_empty
            else 0.0
        )

        if (
            first_ratio <= 0.25
            and second_ratio > first_ratio
        ):
            return True

        # Common contract-table headings.
        header_terms = {
            "description",
            "clause",
            "section",
            "item",
            "service",
            "requirement",
            "amount",
            "price",
            "rate",
            "fee",
            "date",
            "deadline",
            "term",
            "obligation",
            "responsibility",
            "party",
            "status",
            "risk",
            "category",
            "metric",
            "target",
            "sla",
        }

        matches = 0

        for value in first_non_empty:
            normalized = value.lower()

            if any(
                term in normalized
                for term in header_terms
            ):
                matches += 1

        return matches > 0

    @staticmethod
    def _looks_numeric(
        value: str,
    ) -> bool:

        value = value.strip()

        if not value:
            return False

        value = re.sub(
            r"[$€£₹,%]",
            "",
            value,
        )

        value = value.replace(
            ",",
            "",
        )

        try:
            float(value)
            return True

        except ValueError:
            return False

    # ==============================================================
    # Row normalization
    # ==============================================================

    def _normalize_rows(
        self,
        rows: Iterable[
            Iterable[Any]
        ],
    ) -> list[list[str]]:

        normalized: list[
            list[str]
        ] = []

        for row in rows:
            values = [
                self._normalize_cell(
                    value
                )
                for value in row
            ]

            if (
                self.config.remove_empty_rows
                and not any(
                    value.strip()
                    for value in values
                )
            ):
                continue

            normalized.append(
                values
            )

        if (
            not normalized
            or not self.config
            .remove_empty_columns
        ):
            return normalized

        max_columns = max(
            len(row)
            for row in normalized
        )

        keep_columns: list[int] = []

        for column in range(
            max_columns
        ):
            if any(
                column < len(row)
                and row[column].strip()
                for row in normalized
            ):
                keep_columns.append(
                    column
                )

        cleaned: list[
            list[str]
        ] = []

        for row in normalized:
            cleaned.append(
                [
                    row[column]
                    if column < len(row)
                    else ""
                    for column
                    in keep_columns
                ]
            )

        return cleaned

    def _normalize_cell(
        self,
        value: Any,
    ) -> str:

        if value is None:
            return ""

        text = str(value)

        text = text.replace(
            "\x00",
            "",
        )

        if self.config.normalize_whitespace:
            text = re.sub(
                r"[ \t]+",
                " ",
                text,
            )

            text = re.sub(
                r"\n{3,}",
                "\n\n",
                text,
            )

        return text.strip()

    # ==============================================================
    # Result construction
    # ==============================================================

    def _build_result(
        self,
        tables: list[ExtractedTable],
        *,
        filename: Optional[str],
        source_type: TableSourceType,
        processed_pages: int,
        failed_pages: int,
        warnings: list[str],
        metadata: dict[str, Any],
    ) -> TableExtractionResult:

        if not tables:
            if failed_pages:
                status = (
                    TableExtractionStatus.FAILED
                )
            else:
                status = (
                    TableExtractionStatus.NO_TABLES
                )

        elif failed_pages:
            status = (
                TableExtractionStatus.PARTIAL
            )

        else:
            status = (
                TableExtractionStatus.SUCCESS
            )

        return TableExtractionResult(
            status=status,
            tables=tables,
            filename=filename,
            source_type=source_type,
            processed_pages=processed_pages,
            failed_pages=failed_pages,
            warnings=warnings,
            metadata={
                **metadata,
                "total_rows": sum(
                    table.row_count
                    for table in tables
                ),
                "total_cells": sum(
                    len(table.cells)
                    for table in tables
                ),
            },
        )

    # ==============================================================
    # Input handling
    # ==============================================================

    def _load_content(
        self,
        source: (
            str
            | Path
            | bytes
            | bytearray
        ),
    ) -> bytes:

        if isinstance(
            source,
            (bytes, bytearray),
        ):
            return bytes(source)

        path = Path(source)

        if not path.exists():
            raise TableExtractionError(
                f"File does not exist: {path}",
                code="FILE_NOT_FOUND",
            )

        if not path.is_file():
            raise TableExtractionError(
                f"Source is not a file: {path}",
                code="INVALID_SOURCE",
            )

        try:
            return path.read_bytes()

        except OSError as exc:
            raise TableExtractionError(
                f"Unable to read file: {path}",
                code="FILE_READ_ERROR",
                cause=exc,
            ) from exc

    def _validate_content(
        self,
        content: bytes,
    ) -> None:

        if not content:
            raise TableExtractionError(
                "Document is empty.",
                code="EMPTY_DOCUMENT",
            )

        if (
            len(content)
            > self.config.max_document_bytes
        ):
            raise TableExtractionError(
                "Document exceeds configured "
                "table-extraction size limit.",
                code="DOCUMENT_TOO_LARGE",
            )

    @staticmethod
    def _resolve_filename(
        source: Any,
        filename: Optional[str],
    ) -> Optional[str]:

        if filename:
            return Path(filename).name

        if isinstance(
            source,
            (str, Path),
        ):
            return Path(source).name

        return None

    def _resolve_source_type(
        self,
        filename: Optional[str],
        source_type: Optional[
            TableSourceType | str
        ],
    ) -> TableSourceType:

        if isinstance(
            source_type,
            TableSourceType,
        ):
            return source_type

        if isinstance(
            source_type,
            str,
        ):
            try:
                return TableSourceType(
                    source_type.lower()
                )
            except ValueError:
                pass

        if not filename:
            return (
                TableSourceType.UNKNOWN
            )

        extension = (
            Path(filename)
            .suffix
            .lower()
        )

        return self.EXTENSION_MAP.get(
            extension,
            TableSourceType.UNKNOWN,
        )

    # ==============================================================
    # Limits
    # ==============================================================

    def _table_limit_reached(
        self,
        tables: Sequence[
            ExtractedTable
        ],
    ) -> bool:

        limit = (
            self.config.max_tables
        )

        return (
            limit is not None
            and len(tables) >= limit
        )

    def _apply_table_limit(
        self,
        tables: list[
            ExtractedTable
        ],
    ) -> list[ExtractedTable]:

        limit = (
            self.config.max_tables
        )

        if limit is None:
            return tables

        return tables[:limit]

    # ==============================================================
    # Utility
    # ==============================================================

    @staticmethod
    def _bbox(
        value: Any,
    ) -> Optional[
        tuple[
            float,
            float,
            float,
            float,
        ]
    ]:

        if value is None:
            return None

        try:
            if hasattr(value, "x0"):
                return (
                    float(value.x0),
                    float(value.y0),
                    float(value.x1),
                    float(value.y1),
                )

            if len(value) == 4:
                return tuple(
                    float(item)
                    for item in value
                )

        except Exception:
            return None

        return None

    @staticmethod
    def _disabled_result(
        source_type: TableSourceType,
        filename: Optional[str],
    ) -> TableExtractionResult:

        return TableExtractionResult(
            status=(
                TableExtractionStatus.SKIPPED
            ),
            filename=filename,
            source_type=source_type,
            warnings=[
                f"{source_type.value.upper()} "
                "table extraction is disabled."
            ],
        )


def extract_tables(
    source: (
        str
        | Path
        | bytes
        | bytearray
    ),
    *,
    filename: Optional[str] = None,
    source_type: Optional[
        TableSourceType | str
    ] = None,
) -> TableExtractionResult:
    """
    Convenience API used by IngestionPipeline.
    """

    extractor = TableExtractor()

    return extractor.extract(
        source,
        filename=filename,
        source_type=source_type,
    )
"""
DOCX document processor.

Path:
    backend/ingestion/docx_processor.py

Responsibilities:
    - Parse Microsoft Word .docx files.
    - Extract paragraphs and structural information.
    - Extract tables.
    - Extract headers and footers.
    - Preserve hyperlinks where possible.
    - Preserve paragraph/style/list metadata.
    - Extract document core properties.
    - Detect explicit page breaks.
    - Return a parser-friendly structured representation.

Non-responsibilities:
    - OCR.
    - Legal classification.
    - Clause/risk analysis.
    - Chunking.
    - Embeddings.
    - RAG.
    - Legal interpretation.

The processor is intentionally conservative because legal formatting
often carries semantic meaning.
"""

from __future__ import annotations

from dataclasses import dataclass, field, asdict
from enum import Enum
from io import BytesIO
from typing import Any, Optional
import logging
import re
import zipfile
import xml.etree.ElementTree as ET


logger = logging.getLogger(__name__)


# ======================================================================
# Constants
# ======================================================================

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

NS = {
    "w": W_NS,
    "r": R_NS,
    "pr": REL_NS,
}


# ======================================================================
# Enums
# ======================================================================


class BlockType(str, Enum):
    PARAGRAPH = "paragraph"
    HEADING = "heading"
    TABLE = "table"
    LIST_ITEM = "list_item"
    PAGE_BREAK = "page_break"


class Alignment(str, Enum):
    LEFT = "left"
    CENTER = "center"
    RIGHT = "right"
    JUSTIFY = "justify"
    UNKNOWN = "unknown"


# ======================================================================
# Data models
# ======================================================================


@dataclass(slots=True)
class DocxRun:
    """A run of text within a DOCX paragraph."""

    text: str

    bold: bool = False
    italic: bool = False
    underline: bool = False
    strike: bool = False

    style: Optional[str] = None

    metadata: dict[str, Any] = field(
        default_factory=dict
    )


@dataclass(slots=True)
class DocxParagraph:
    """Structured DOCX paragraph."""

    index: int

    text: str

    style: Optional[str] = None

    style_id: Optional[str] = None

    heading_level: Optional[int] = None

    alignment: Alignment = Alignment.UNKNOWN

    is_list_item: bool = False

    list_level: Optional[int] = None

    numbering_id: Optional[str] = None

    page_break_before: bool = False

    contains_page_break: bool = False

    runs: list[DocxRun] = field(
        default_factory=list
    )

    hyperlinks: list[str] = field(
        default_factory=list
    )

    metadata: dict[str, Any] = field(
        default_factory=dict
    )


@dataclass(slots=True)
class DocxTableCell:
    """DOCX table cell."""

    row_index: int

    column_index: int

    text: str

    paragraphs: list[str] = field(
        default_factory=list
    )

    metadata: dict[str, Any] = field(
        default_factory=dict
    )


@dataclass(slots=True)
class DocxTable:
    """DOCX table."""

    index: int

    rows: list[
        list[DocxTableCell]
    ] = field(
        default_factory=list
    )

    headers: list[str] = field(
        default_factory=list
    )

    metadata: dict[str, Any] = field(
        default_factory=dict
    )


@dataclass(slots=True)
class DocxHeaderFooter:
    """DOCX header/footer."""

    section_index: int

    kind: str

    text: str

    paragraphs: list[str] = field(
        default_factory=list
    )


@dataclass(slots=True)
class DocxDocumentProperties:
    """DOCX core properties."""

    title: Optional[str] = None
    subject: Optional[str] = None
    creator: Optional[str] = None
    last_modified_by: Optional[str] = None
    description: Optional[str] = None
    keywords: Optional[str] = None
    category: Optional[str] = None
    created: Optional[str] = None
    modified: Optional[str] = None
    revision: Optional[str] = None


@dataclass(slots=True)
class DocxProcessingResult:
    """Complete DOCX processing result."""

    filename: str

    text: str

    paragraphs: list[DocxParagraph] = field(
        default_factory=list
    )

    tables: list[DocxTable] = field(
        default_factory=list
    )

    headers_footers: list[
        DocxHeaderFooter
    ] = field(
        default_factory=list
    )

    properties: DocxDocumentProperties = field(
        default_factory=DocxDocumentProperties
    )

    metadata: dict[str, Any] = field(
        default_factory=dict
    )

    warnings: list[str] = field(
        default_factory=list
    )

    page_count: int = 1

    paragraph_count: int = 0

    table_count: int = 0

    character_count: int = 0

    word_count: int = 0


@dataclass(slots=True)
class DocxProcessorConfig:
    """Processor configuration."""

    extract_tables: bool = True

    extract_headers_footers: bool = True

    extract_properties: bool = True

    extract_hyperlinks: bool = True

    preserve_runs: bool = True

    preserve_empty_paragraphs: bool = False

    detect_page_breaks: bool = True

    max_document_size_bytes: int = 50 * 1024 * 1024

    include_headers_in_text: bool = True

    include_footers_in_text: bool = True

    include_tables_in_text: bool = True


# ======================================================================
# Processor
# ======================================================================


class DocxProcessor:
    """
    Microsoft Word DOCX processor.

    Uses python-docx for normal document extraction and performs
    lightweight XML inspection for information that is not always
    exposed conveniently by the high-level API.
    """

    def __init__(
        self,
        config: Optional[
            DocxProcessorConfig
        ] = None,
    ) -> None:
        self.config = (
            config
            or DocxProcessorConfig()
        )

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def process(
        self,
        content: bytes,
        *,
        filename: str = "document.docx",
    ) -> DocxProcessingResult:
        """
        Process DOCX bytes.

        Args:
            content: Raw .docx bytes.
            filename: Original filename.

        Returns:
            DocxProcessingResult
        """

        self._validate(
            content,
            filename,
        )

        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError(
                "python-docx is required for DOCX processing."
            ) from exc

        document = Document(
            BytesIO(content)
        )

        result = DocxProcessingResult(
            filename=filename,
        )

        if self.config.extract_properties:
            result.properties = (
                self._extract_properties(
                    document
                )
            )

        result.paragraphs = (
            self._extract_paragraphs(
                document
            )
        )

        if self.config.extract_tables:
            result.tables = (
                self._extract_tables(
                    document
                )
            )

        if self.config.extract_headers_footers:
            result.headers_footers = (
                self._extract_headers_footers(
                    document
                )
            )

        if self.config.detect_page_breaks:
            self._apply_xml_page_break_detection(
                content,
                result,
            )

        result.text = (
            self._build_document_text(
                result
            )
        )

        result.paragraph_count = len(
            result.paragraphs
        )

        result.table_count = len(
            result.tables
        )

        result.character_count = len(
            result.text
        )

        result.word_count = len(
            re.findall(
                r"\S+",
                result.text,
            )
        )

        result.page_count = (
            self._estimate_page_count(
                result
            )
        )

        result.metadata = {
            "format": "docx",
            "content_type": (
                "application/vnd.openxmlformats-"
                "officedocument.wordprocessingml.document"
            ),
            "paragraph_count": (
                result.paragraph_count
            ),
            "table_count": (
                result.table_count
            ),
            "page_count": (
                result.page_count
            ),
        }

        return result

    def process_file(
        self,
        file_path: str,
    ) -> DocxProcessingResult:
        """
        Process a DOCX file from disk.
        """

        from pathlib import Path

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                str(path)
            )

        content = path.read_bytes()

        return self.process(
            content,
            filename=path.name,
        )

    # ------------------------------------------------------------------
    # Validation
    # ------------------------------------------------------------------

    def _validate(
        self,
        content: bytes,
        filename: str,
    ) -> None:
        if not isinstance(
            content,
            bytes,
        ):
            raise TypeError(
                "DOCX content must be bytes."
            )

        if not content:
            raise ValueError(
                "DOCX document is empty."
            )

        if len(content) > (
            self.config.max_document_size_bytes
        ):
            raise ValueError(
                (
                    f"DOCX document exceeds "
                    f"{self.config.max_document_size_bytes} "
                    "byte limit."
                )
            )

        if not self._is_docx_package(
            content
        ):
            raise ValueError(
                f"{filename} is not a valid DOCX package."
            )

    @staticmethod
    def _is_docx_package(
        content: bytes,
    ) -> bool:
        """
        DOCX is an OOXML ZIP package.
        """

        if not content.startswith(
            b"PK"
        ):
            return False

        try:
            with zipfile.ZipFile(
                BytesIO(content)
            ) as archive:
                names = set(
                    archive.namelist()
                )

                return (
                    "word/document.xml"
                    in names
                    and "[Content_Types].xml"
                    in names
                )

        except zipfile.BadZipFile:
            return False

    # ------------------------------------------------------------------
    # Paragraph extraction
    # ------------------------------------------------------------------

    def _extract_paragraphs(
        self,
        document: Any,
    ) -> list[DocxParagraph]:
        paragraphs: list[
            DocxParagraph
        ] = []

        for index, paragraph in enumerate(
            document.paragraphs
        ):
            text = paragraph.text or ""

            if (
                not text.strip()
                and not self.config.preserve_empty_paragraphs
            ):
                continue

            style_name = None
            style_id = None

            try:
                style_name = (
                    paragraph.style.name
                    if paragraph.style
                    else None
                )

                style_id = (
                    paragraph.style.style_id
                    if paragraph.style
                    else None
                )
            except Exception:
                pass

            heading_level = (
                self._get_heading_level(
                    paragraph,
                    style_name,
                )
            )

            is_list_item, list_level, numbering_id = (
                self._get_list_metadata(
                    paragraph
                )
            )

            alignment = (
                self._get_alignment(
                    paragraph
                )
            )

            page_break_before = (
                self._has_page_break_before(
                    paragraph
                )
            )

            runs = (
                self._extract_runs(
                    paragraph
                )
                if self.config.preserve_runs
                else []
            )

            hyperlinks = (
                self._extract_hyperlinks(
                    paragraph
                )
                if self.config.extract_hyperlinks
                else []
            )

            paragraphs.append(
                DocxParagraph(
                    index=index,
                    text=text,
                    style=style_name,
                    style_id=style_id,
                    heading_level=heading_level,
                    alignment=alignment,
                    is_list_item=is_list_item,
                    list_level=list_level,
                    numbering_id=numbering_id,
                    page_break_before=page_break_before,
                    runs=runs,
                    hyperlinks=hyperlinks,
                )
            )

        return paragraphs

    # ------------------------------------------------------------------
    # Runs
    # ------------------------------------------------------------------

    def _extract_runs(
        self,
        paragraph: Any,
    ) -> list[DocxRun]:
        runs: list[DocxRun] = []

        for run in paragraph.runs:
            text = run.text or ""

            if not text:
                continue

            style = None

            try:
                if run.style:
                    style = run.style.name
            except Exception:
                pass

            runs.append(
                DocxRun(
                    text=text,
                    bold=bool(
                        run.bold
                    ),
                    italic=bool(
                        run.italic
                    ),
                    underline=bool(
                        run.underline
                    ),
                    strike=bool(
                        run.font.strike
                    ),
                    style=style,
                )
            )

        return runs

    # ------------------------------------------------------------------
    # Headings
    # ------------------------------------------------------------------

    def _get_heading_level(
        self,
        paragraph: Any,
        style_name: Optional[str],
    ) -> Optional[int]:
        """
        Detect Word Heading 1..9 styles.

        Also handles common variants such as:
            Heading 1
            heading 2
            Heading1
        """

        if not style_name:
            return None

        match = re.search(
            r"heading\s*([1-9])",
            style_name,
            re.IGNORECASE,
        )

        if match:
            return int(
                match.group(1)
            )

        # Outline level fallback.
        try:
            ppr = paragraph._p.pPr

            if ppr is not None:
                outline = ppr.find(
                    f"{{{W_NS}}}outlineLvl"
                )

                if outline is not None:
                    value = outline.get(
                        f"{{{W_NS}}}val"
                    )

                    if value is not None:
                        level = int(value) + 1

                        if 1 <= level <= 9:
                            return level
        except Exception:
            pass

        return None

    # ------------------------------------------------------------------
    # Lists
    # ------------------------------------------------------------------

    def _get_list_metadata(
        self,
        paragraph: Any,
    ) -> tuple[
        bool,
        Optional[int],
        Optional[str],
    ]:
        """
        Extract numbering metadata from paragraph XML.
        """

        try:
            ppr = paragraph._p.pPr

            if ppr is None:
                return False, None, None

            num_pr = ppr.find(
                f"{{{W_NS}}}numPr"
            )

            if num_pr is None:
                return False, None, None

            ilvl = num_pr.find(
                f"{{{W_NS}}}ilvl"
            )

            num_id = num_pr.find(
                f"{{{W_NS}}}numId"
            )

            level = None
            numbering_id = None

            if ilvl is not None:
                raw_level = ilvl.get(
                    f"{{{W_NS}}}val"
                )

                if raw_level is not None:
                    level = int(
                        raw_level
                    )

            if num_id is not None:
                numbering_id = (
                    num_id.get(
                        f"{{{W_NS}}}val"
                    )
                )

            return (
                True,
                level,
                numbering_id,
            )

        except Exception:
            return False, None, None

    # ------------------------------------------------------------------
    # Alignment
    # ------------------------------------------------------------------

    @staticmethod
    def _get_alignment(
        paragraph: Any,
    ) -> Alignment:
        try:
            value = paragraph.alignment

            if value is None:
                return Alignment.UNKNOWN

            name = getattr(
                value,
                "name",
                "",
            ).lower()

            mapping = {
                "left": Alignment.LEFT,
                "center": Alignment.CENTER,
                "right": Alignment.RIGHT,
                "justify": Alignment.JUSTIFY,
            }

            return mapping.get(
                name,
                Alignment.UNKNOWN,
            )

        except Exception:
            return Alignment.UNKNOWN

    # ------------------------------------------------------------------
    # Page breaks
    # ------------------------------------------------------------------

    @staticmethod
    def _has_page_break_before(
        paragraph: Any,
    ) -> bool:
        try:
            ppr = paragraph._p.pPr

            if ppr is None:
                return False

            page_break = ppr.find(
                f"{{{W_NS}}}pageBreakBefore"
            )

            return page_break is not None

        except Exception:
            return False

    def _apply_xml_page_break_detection(
        self,
        content: bytes,
        result: DocxProcessingResult,
    ) -> None:
        """
        Detect explicit page breaks from document.xml.

        This is intentionally not presented as an exact rendered
        page count. Word pagination depends on fonts, layout,
        printer settings, margins, images, and rendering engine.
        """

        try:
            with zipfile.ZipFile(
                BytesIO(content)
            ) as archive:
                xml = archive.read(
                    "word/document.xml"
                )

            root = ET.fromstring(
                xml
            )

            paragraphs_xml = root.findall(
                ".//w:body/w:p",
                NS,
            )

            for index, p_xml in enumerate(
                paragraphs_xml
            ):
                if index >= len(
                    result.paragraphs
                ):
                    break

                breaks = p_xml.findall(
                    ".//w:br",
                    NS,
                )

                has_break = any(
                    br.get(
                        f"{{{W_NS}}}type"
                    )
                    == "page"
                    for br in breaks
                )

                if has_break:
                    result.paragraphs[
                        index
                    ].contains_page_break = True

        except Exception as exc:
            result.warnings.append(
                "Unable to inspect explicit page breaks."
            )

            logger.debug(
                "Page-break detection failed: %s",
                exc,
            )

    # ------------------------------------------------------------------
    # Hyperlinks
    # ------------------------------------------------------------------

    def _extract_hyperlinks(
        self,
        paragraph: Any,
    ) -> list[str]:
        """
        Extract hyperlink URLs from paragraph XML relationships.
        """

        links: list[str] = []

        try:
            part = paragraph.part

            for hyperlink in paragraph._p.findall(
                f".//{{{W_NS}}}hyperlink"
            ):
                relationship_id = (
                    hyperlink.get(
                        f"{{{R_NS}}}id"
                    )
                )

                if not relationship_id:
                    continue

                try:
                    target = (
                        part.rels[
                            relationship_id
                        ].target_ref
                    )

                    if target:
                        links.append(
                            str(target)
                        )
                except Exception:
                    continue

        except Exception:
            pass

        return list(
            dict.fromkeys(links)
        )

    # ------------------------------------------------------------------
    # Tables
    # ------------------------------------------------------------------

    def _extract_tables(
        self,
        document: Any,
    ) -> list[DocxTable]:
        tables: list[
            DocxTable
        ] = []

        for table_index, table in enumerate(
            document.tables
        ):
            rows: list[
                list[DocxTableCell]
            ] = []

            for row_index, row in enumerate(
                table.rows
            ):
                cells: list[
                    DocxTableCell
                ] = []

                for column_index, cell in enumerate(
                    row.cells
                ):
                    paragraphs = [
                        paragraph.text
                        for paragraph in (
                            cell.paragraphs
                        )
                    ]

                    cell_text = "\n".join(
                        text
                        for text in paragraphs
                        if text
                    )

                    cells.append(
                        DocxTableCell(
                            row_index=row_index,
                            column_index=column_index,
                            text=cell_text,
                            paragraphs=paragraphs,
                        )
                    )

                rows.append(cells)

            headers = []

            if rows:
                headers = [
                    cell.text
                    for cell in rows[0]
                ]

            tables.append(
                DocxTable(
                    index=table_index,
                    rows=rows,
                    headers=headers,
                    metadata={
                        "row_count": len(rows),
                        "column_count": (
                            max(
                                (
                                    len(row)
                                    for row in rows
                                ),
                                default=0,
                            )
                        ),
                    },
                )
            )

        return tables

    # ------------------------------------------------------------------
    # Headers / footers
    # ------------------------------------------------------------------

    def _extract_headers_footers(
        self,
        document: Any,
    ) -> list[DocxHeaderFooter]:
        result: list[
            DocxHeaderFooter
        ] = []

        for section_index, section in enumerate(
            document.sections
        ):
            header = section.header

            header_paragraphs = [
                paragraph.text
                for paragraph in (
                    header.paragraphs
                )
                if paragraph.text.strip()
            ]

            if header_paragraphs:
                result.append(
                    DocxHeaderFooter(
                        section_index=section_index,
                        kind="header",
                        text="\n".join(
                            header_paragraphs
                        ),
                        paragraphs=header_paragraphs,
                    )
                )

            footer = section.footer

            footer_paragraphs = [
                paragraph.text
                for paragraph in (
                    footer.paragraphs
                )
                if paragraph.text.strip()
            ]

            if footer_paragraphs:
                result.append(
                    DocxHeaderFooter(
                        section_index=section_index,
                        kind="footer",
                        text="\n".join(
                            footer_paragraphs
                        ),
                        paragraphs=footer_paragraphs,
                    )
                )

        return result

    # ------------------------------------------------------------------
    # Properties
    # ------------------------------------------------------------------

    @staticmethod
    def _extract_properties(
        document: Any,
    ) -> DocxDocumentProperties:
        try:
            properties = document.core_properties

            return DocxDocumentProperties(
                title=properties.title,
                subject=properties.subject,
                creator=properties.author,
                last_modified_by=(
                    properties.last_modified_by
                ),
                description=(
                    properties.comments
                ),
                keywords=properties.keywords,
                category=properties.category,
                created=(
                    properties.created.isoformat()
                    if properties.created
                    else None
                ),
                modified=(
                    properties.modified.isoformat()
                    if properties.modified
                    else None
                ),
                revision=(
                    str(properties.revision)
                    if properties.revision
                    else None
                ),
            )

        except Exception as exc:
            logger.debug(
                "Unable to extract DOCX properties: %s",
                exc,
            )

            return DocxDocumentProperties()

    # ------------------------------------------------------------------
    # Text construction
    # ------------------------------------------------------------------

    def _build_document_text(
        self,
        result: DocxProcessingResult,
    ) -> str:
        """
        Build canonical text while preserving important legal
        structure.

        Tables are separated from normal paragraphs rather than
        flattened into one unstructured stream.
        """

        blocks: list[str] = []

        if (
            self.config.include_headers_in_text
        ):
            for item in result.headers_footers:
                if item.kind == "header":
                    blocks.append(
                        f"[HEADER]\n{item.text}"
                    )

        for paragraph in result.paragraphs:
            text = paragraph.text

            if not text.strip():
                if (
                    self.config.preserve_empty_paragraphs
                ):
                    blocks.append("")
                continue

            if paragraph.heading_level:
                blocks.append(
                    (
                        f"[HEADING "
                        f"{paragraph.heading_level}] "
                        f"{text}"
                    )
                )

            elif paragraph.is_list_item:
                prefix = self._list_prefix(
                    paragraph
                )

                blocks.append(
                    f"{prefix}{text}"
                )

            else:
                blocks.append(
                    text
                )

            if paragraph.contains_page_break:
                blocks.append(
                    "[PAGE_BREAK]"
                )

        if (
            self.config.include_tables_in_text
        ):
            for table in result.tables:
                blocks.append(
                    self._table_to_text(
                        table
                    )
                )

        if (
            self.config.include_footers_in_text
        ):
            for item in result.headers_footers:
                if item.kind == "footer":
                    blocks.append(
                        f"[FOOTER]\n{item.text}"
                    )

        return "\n\n".join(
            block
            for block in blocks
            if block is not None
        ).strip()

    @staticmethod
    def _list_prefix(
        paragraph: DocxParagraph,
    ) -> str:
        level = (
            paragraph.list_level
            or 0
        )

        indentation = (
            "  " * level
        )

        return (
            f"{indentation}- "
        )

    @staticmethod
    def _table_to_text(
        table: DocxTable,
    ) -> str:
        """
        Convert a table into a stable text representation.

        This is for downstream indexing/search only.
        Original table structure remains available in `tables`.
        """

        lines = [
            f"[TABLE {table.index + 1}]"
        ]

        for row in table.rows:
            values = [
                cell.text.replace(
                    "\n",
                    " | ",
                )
                for cell in row
            ]

            lines.append(
                " | ".join(values)
            )

        return "\n".join(
            lines
        )

    # ------------------------------------------------------------------
    # Page estimation
    # ------------------------------------------------------------------

    @staticmethod
    def _estimate_page_count(
        result: DocxProcessingResult,
    ) -> int:
        """
        Estimate page count.

        Exact pagination requires a Word-compatible rendering engine.
        We therefore use explicit page breaks as the reliable signal
        and default to one page for documents without them.
        """

        explicit_breaks = sum(
            1
            for paragraph in result.paragraphs
            if paragraph.contains_page_break
            or paragraph.page_break_before
        )

        return max(
            1,
            explicit_breaks + 1,
        )

    # ------------------------------------------------------------------
    # Serialization
    # ------------------------------------------------------------------

    @staticmethod
    def to_dict(
        result: DocxProcessingResult,
    ) -> dict[str, Any]:
        """
        Convert processing result into JSON-compatible dictionary.
        """

        return asdict(
            result
        )


# ======================================================================
# Convenience function
# ======================================================================


_default_processor = DocxProcessor()


def process_docx(
    content: bytes,
    *,
    filename: str = "document.docx",
) -> DocxProcessingResult:
    """
    Convenience API.
    """

    return _default_processor.process(
        content,
        filename=filename,
    )


__all__ = [
    "BlockType",
    "Alignment",
    "DocxRun",
    "DocxParagraph",
    "DocxTableCell",
    "DocxTable",
    "DocxHeaderFooter",
    "DocxDocumentProperties",
    "DocxProcessingResult",
    "DocxProcessorConfig",
    "DocxProcessor",
    "process_docx",
]
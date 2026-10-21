# backend/ingestion/page_extractor.py

from __future__ import annotations

from dataclasses import dataclass, field
from enum import Enum
from io import BytesIO
from pathlib import Path
from typing import Any, Optional
import re


class PageExtractionStatus(str, Enum):
    SUCCESS = "success"
    PARTIAL = "partial"
    NO_PAGES = "no_pages"
    FAILED = "failed"


class PageContentType(str, Enum):
    TEXT = "text"
    IMAGE = "image"
    MIXED = "mixed"
    EMPTY = "empty"
    UNKNOWN = "unknown"


@dataclass(slots=True)
class PageResult:
    """Normalized representation of one document page."""

    page_number: int
    text: str = ""

    content_type: PageContentType = PageContentType.UNKNOWN

    character_count: int = 0
    word_count: int = 0

    has_text: bool = False
    has_images: bool = False
    image_count: int = 0

    width: Optional[float] = None
    height: Optional[float] = None

    rotation: Optional[int] = None

    metadata: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)

    def __post_init__(self) -> None:
        self.text = self.text or ""

        cleaned = self.text.strip()

        self.character_count = len(cleaned)
        self.word_count = len(cleaned.split())
        self.has_text = bool(cleaned)

        if self.has_text and self.has_images:
            self.content_type = PageContentType.MIXED
        elif self.has_text:
            self.content_type = PageContentType.TEXT
        elif self.has_images:
            self.content_type = PageContentType.IMAGE
        else:
            self.content_type = PageContentType.EMPTY

    def to_dict(self) -> dict[str, Any]:
        return {
            "page_number": self.page_number,
            "text": self.text,
            "content_type": self.content_type.value,
            "character_count": self.character_count,
            "word_count": self.word_count,
            "has_text": self.has_text,
            "has_images": self.has_images,
            "image_count": self.image_count,
            "width": self.width,
            "height": self.height,
            "rotation": self.rotation,
            "metadata": self.metadata,
            "warnings": self.warnings,
        }


@dataclass(slots=True)
class PageExtractionResult:
    """Result returned by the page extraction layer."""

    status: PageExtractionStatus

    pages: list[PageResult] = field(default_factory=list)

    page_count: int = 0
    extracted_pages: int = 0
    failed_pages: int = 0

    document_type: Optional[str] = None
    filename: Optional[str] = None

    has_text: bool = False
    has_images: bool = False
    is_scanned: bool = False

    warnings: list[str] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        return {
            "status": self.status.value,
            "pages": [
                page.to_dict()
                for page in self.pages
            ],
            "page_count": self.page_count,
            "extracted_pages": self.extracted_pages,
            "failed_pages": self.failed_pages,
            "document_type": self.document_type,
            "filename": self.filename,
            "has_text": self.has_text,
            "has_images": self.has_images,
            "is_scanned": self.is_scanned,
            "warnings": self.warnings,
            "metadata": self.metadata,
        }


@dataclass(slots=True)
class PageExtractorConfig:
    """
    Configuration for page extraction.

    Extraction is intentionally conservative because page boundaries
    are important for legal citations and downstream RAG provenance.
    """

    preserve_page_boundaries: bool = True

    normalize_whitespace: bool = True

    detect_images: bool = True

    detect_rotation: bool = True

    max_pages: Optional[int] = None

    min_text_length: int = 2

    max_document_bytes: int = 100 * 1024 * 1024

    inspect_pdf: bool = True

    inspect_docx: bool = True

    inspect_pptx: bool = True

    inspect_xlsx: bool = True


class PageExtractionError(Exception):
    """Raised for fatal page extraction failures."""

    def __init__(
        self,
        message: str,
        *,
        code: str = "PAGE_EXTRACTION_ERROR",
        cause: Optional[Exception] = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.cause = cause


class PageExtractor:
    """
    Extract page-level content from supported document formats.

    Responsibilities:
        - preserve page boundaries
        - extract page-level text
        - detect images
        - expose page dimensions/rotation where available
        - provide page provenance for RAG/citations
        - identify likely scanned pages

    Non-responsibilities:
        - OCR itself
        - document normalization
        - language detection
        - legal interpretation
        - clause/risk classification
    """

    PDF_EXTENSIONS = {".pdf"}

    DOCX_EXTENSIONS = {".docx"}

    PPTX_EXTENSIONS = {".pptx"}

    XLSX_EXTENSIONS = {".xlsx"}

    IMAGE_EXTENSIONS = {
        ".png",
        ".jpg",
        ".jpeg",
        ".tif",
        ".tiff",
        ".bmp",
        ".webp",
    }

    def __init__(
        self,
        config: Optional[PageExtractorConfig] = None,
    ) -> None:
        self.config = config or PageExtractorConfig()

    # ================================================================
    # Public API
    # ================================================================

    def extract(
        self,
        source: str | Path | bytes | bytearray,
        *,
        filename: Optional[str] = None,
        document_type: Optional[str] = None,
    ) -> PageExtractionResult:
        """
        Extract pages from a file path or raw bytes.
        """

        content = self._load_content(source)

        self._validate_content(content)

        resolved_filename = self._resolve_filename(
            source,
            filename,
        )

        extension = self._extension(
            resolved_filename
        )

        normalized_type = (
            document_type.lower().strip()
            if document_type
            else None
        )

        try:
            if (
                extension in self.PDF_EXTENSIONS
                or normalized_type == "pdf"
            ):
                return self.extract_pdf(
                    content,
                    filename=resolved_filename,
                )

            if (
                extension in self.DOCX_EXTENSIONS
                or normalized_type == "docx"
            ):
                return self.extract_docx(
                    content,
                    filename=resolved_filename,
                )

            if (
                extension in self.PPTX_EXTENSIONS
                or normalized_type == "pptx"
            ):
                return self.extract_pptx(
                    content,
                    filename=resolved_filename,
                )

            if (
                extension in self.XLSX_EXTENSIONS
                or normalized_type == "xlsx"
            ):
                return self.extract_xlsx(
                    content,
                    filename=resolved_filename,
                )

            if extension in self.IMAGE_EXTENSIONS:
                return self.extract_image(
                    content,
                    filename=resolved_filename,
                )

            return PageExtractionResult(
                status=PageExtractionStatus.NO_PAGES,
                document_type=extension.lstrip(".")
                if extension
                else "unknown",
                filename=resolved_filename,
                warnings=[
                    "Unsupported document format for page extraction."
                ],
            )

        except PageExtractionError:
            raise

        except Exception as exc:
            raise PageExtractionError(
                "Unexpected page extraction failure.",
                code="UNEXPECTED_ERROR",
                cause=exc,
            ) from exc

    def extract_pdf(
        self,
        content: bytes,
        *,
        filename: Optional[str] = None,
    ) -> PageExtractionResult:
        if not self.config.inspect_pdf:
            return self._disabled_result(
                "pdf",
                filename,
            )

        try:
            import fitz
        except ImportError as exc:
            raise PageExtractionError(
                "PyMuPDF is required for PDF page extraction.",
                code="PYMUPDF_NOT_INSTALLED",
                cause=exc,
            ) from exc

        document = None

        try:
            document = fitz.open(
                stream=content,
                filetype="pdf",
            )

            if document.needs_pass:
                raise PageExtractionError(
                    "Encrypted PDF requires a password.",
                    code="ENCRYPTED_DOCUMENT",
                )

            total_pages = document.page_count

            page_limit = self.config.max_pages

            if page_limit is None:
                page_limit = total_pages
            else:
                page_limit = min(
                    page_limit,
                    total_pages,
                )

            pages: list[PageResult] = []
            warnings: list[str] = []
            failed_pages = 0

            for index in range(page_limit):
                page_number = index + 1

                try:
                    page = document.load_page(index)

                    result = self._extract_pdf_page(
                        page,
                        page_number,
                    )

                    pages.append(result)

                except Exception as exc:
                    failed_pages += 1

                    warnings.append(
                        f"Failed to extract page "
                        f"{page_number}: {exc}"
                    )

            return self._build_result(
                pages,
                total_pages=total_pages,
                failed_pages=failed_pages,
                document_type="pdf",
                filename=filename,
                warnings=warnings,
            )

        except PageExtractionError:
            raise

        except Exception as exc:
            raise PageExtractionError(
                "Failed to open or inspect PDF.",
                code="PDF_EXTRACTION_ERROR",
                cause=exc,
            ) from exc

        finally:
            if document is not None:
                try:
                    document.close()
                except Exception:
                    pass

    def extract_docx(
        self,
        content: bytes,
        *,
        filename: Optional[str] = None,
    ) -> PageExtractionResult:
        """
        DOCX does not have a reliable native page model.

        python-docx exposes logical document structure rather than
        rendered pages. Therefore this method creates a logical
        page representation rather than pretending that paragraph
        boundaries are physical page boundaries.
        """

        if not self.config.inspect_docx:
            return self._disabled_result(
                "docx",
                filename,
            )

        try:
            from docx import Document
        except ImportError as exc:
            raise PageExtractionError(
                "python-docx is required for DOCX extraction.",
                code="PYTHON_DOCX_NOT_INSTALLED",
                cause=exc,
            ) from exc

        try:
            document = Document(
                BytesIO(content)
            )

            text_parts: list[str] = []

            for paragraph in document.paragraphs:
                text = paragraph.text or ""

                if text.strip():
                    text_parts.append(text)

            # Tables are important in contracts and should not disappear.
            table_text = self._extract_docx_tables(
                document
            )

            if table_text:
                text_parts.append(
                    "\n".join(table_text)
                )

            text = "\n\n".join(
                text_parts
            )

            page = PageResult(
                page_number=1,
                text=self._normalize_text(text),
                has_images=self._docx_has_images(
                    document
                ),
                image_count=self._docx_image_count(
                    document
                ),
                metadata={
                    "logical_page": True,
                    "physical_page_count_available": False,
                    "paragraph_count": len(
                        document.paragraphs
                    ),
                    "table_count": len(
                        document.tables
                    ),
                    "section_count": len(
                        document.sections
                    ),
                },
            )

            warnings = [
                "DOCX does not expose reliable physical page boundaries "
                "through python-docx; page 1 represents the logical "
                "document content."
            ]

            return self._build_result(
                [page],
                total_pages=1,
                failed_pages=0,
                document_type="docx",
                filename=filename,
                warnings=warnings,
            )

        except Exception as exc:
            raise PageExtractionError(
                "Failed to extract DOCX content.",
                code="DOCX_EXTRACTION_ERROR",
                cause=exc,
            ) from exc

    def extract_pptx(
        self,
        content: bytes,
        *,
        filename: Optional[str] = None,
    ) -> PageExtractionResult:
        if not self.config.inspect_pptx:
            return self._disabled_result(
                "pptx",
                filename,
            )

        try:
            from pptx import Presentation
        except ImportError as exc:
            raise PageExtractionError(
                "python-pptx is required for PPTX extraction.",
                code="PYTHON_PPTX_NOT_INSTALLED",
                cause=exc,
            ) from exc

        try:
            presentation = Presentation(
                BytesIO(content)
            )

            pages: list[PageResult] = []

            slides = presentation.slides

            max_slides = self.config.max_pages

            slide_count = len(slides)

            if max_slides is not None:
                slide_count = min(
                    slide_count,
                    max_slides,
                )

            for index in range(slide_count):
                slide = slides[index]

                text_parts: list[str] = []
                image_count = 0

                for shape in slide.shapes:
                    if getattr(
                        shape,
                        "has_text_frame",
                        False,
                    ):
                        text = shape.text or ""

                        if text.strip():
                            text_parts.append(text)

                    if getattr(
                        shape,
                        "shape_type",
                        None,
                    ) == 13:  # PICTURE
                        image_count += 1

                pages.append(
                    PageResult(
                        page_number=index + 1,
                        text=self._normalize_text(
                            "\n".join(text_parts)
                        ),
                        has_images=image_count > 0,
                        image_count=image_count,
                        width=float(
                            presentation.slide_width
                        ),
                        height=float(
                            presentation.slide_height
                        ),
                        metadata={
                            "slide_number": index + 1,
                        },
                    )
                )

            return self._build_result(
                pages,
                total_pages=len(slides),
                failed_pages=0,
                document_type="pptx",
                filename=filename,
                warnings=[],
            )

        except Exception as exc:
            raise PageExtractionError(
                "Failed to extract PPTX slides.",
                code="PPTX_EXTRACTION_ERROR",
                cause=exc,
            ) from exc

    def extract_xlsx(
        self,
        content: bytes,
        *,
        filename: Optional[str] = None,
    ) -> PageExtractionResult:
        if not self.config.inspect_xlsx:
            return self._disabled_result(
                "xlsx",
                filename,
            )

        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise PageExtractionError(
                "openpyxl is required for XLSX extraction.",
                code="OPENPYXL_NOT_INSTALLED",
                cause=exc,
            ) from exc

        workbook = None

        try:
            workbook = load_workbook(
                BytesIO(content),
                read_only=True,
                data_only=True,
            )

            pages: list[PageResult] = []

            sheet_names = workbook.sheetnames

            max_sheets = self.config.max_pages

            if max_sheets is not None:
                sheet_names = sheet_names[
                    :max_sheets
                ]

            for index, sheet_name in enumerate(
                sheet_names,
                start=1,
            ):
                sheet = workbook[sheet_name]

                rows: list[str] = []

                for row in sheet.iter_rows(
                    values_only=True
                ):
                    values = [
                        str(value).strip()
                        for value in row
                        if value is not None
                    ]

                    if values:
                        rows.append(
                            " | ".join(values)
                        )

                pages.append(
                    PageResult(
                        page_number=index,
                        text=self._normalize_text(
                            "\n".join(rows)
                        ),
                        metadata={
                            "sheet_name": sheet_name,
                            "sheet_index": index,
                            "logical_page": True,
                        },
                    )
                )

            warnings = [
                "XLSX worksheets are represented as logical pages; "
                "physical print-page boundaries are not guaranteed."
            ]

            return self._build_result(
                pages,
                total_pages=len(workbook.sheetnames),
                failed_pages=0,
                document_type="xlsx",
                filename=filename,
                warnings=warnings,
            )

        except Exception as exc:
            raise PageExtractionError(
                "Failed to extract XLSX worksheets.",
                code="XLSX_EXTRACTION_ERROR",
                cause=exc,
            ) from exc

        finally:
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass

    def extract_image(
        self,
        content: bytes,
        *,
        filename: Optional[str] = None,
    ) -> PageExtractionResult:
        """
        An image is treated as one logical page.

        OCR remains responsible for extracting text from it.
        """

        try:
            from PIL import Image
        except ImportError as exc:
            raise PageExtractionError(
                "Pillow is required for image inspection.",
                code="PILLOW_NOT_INSTALLED",
                cause=exc,
            ) from exc

        try:
            image = Image.open(
                BytesIO(content)
            )

            width, height = image.size

            page = PageResult(
                page_number=1,
                text="",
                has_images=True,
                image_count=1,
                width=float(width),
                height=float(height),
                metadata={
                    "format": image.format,
                    "mode": image.mode,
                },
            )

            return self._build_result(
                [page],
                total_pages=1,
                failed_pages=0,
                document_type="image",
                filename=filename,
                warnings=[
                    "Image contains no native text; OCR may be required."
                ],
            )

        except Exception as exc:
            raise PageExtractionError(
                "Failed to inspect image.",
                code="IMAGE_EXTRACTION_ERROR",
                cause=exc,
            ) from exc

    # ================================================================
    # PDF page extraction
    # ================================================================

    def _extract_pdf_page(
        self,
        page: Any,
        page_number: int,
    ) -> PageResult:
        text = page.get_text(
            "text"
        ) or ""

        text = self._normalize_text(
            text
        )

        image_count = 0

        if self.config.detect_images:
            try:
                image_list = page.get_images(
                    full=True
                )

                image_count = len(
                    image_list
                )
            except Exception:
                image_count = 0

        rect = page.rect

        rotation = None

        if self.config.detect_rotation:
            try:
                rotation = int(
                    page.rotation
                )
            except Exception:
                rotation = None

        return PageResult(
            page_number=page_number,
            text=text,
            has_images=image_count > 0,
            image_count=image_count,
            width=float(rect.width),
            height=float(rect.height),
            rotation=rotation,
            metadata={
                "pdf_page_index": page_number - 1,
            },
        )

    # ================================================================
    # DOCX helpers
    # ================================================================

    @staticmethod
    def _extract_docx_tables(
        document: Any,
    ) -> list[str]:
        output: list[str] = []

        for table_index, table in enumerate(
            document.tables,
            start=1,
        ):
            rows: list[str] = []

            for row in table.rows:
                cells = [
                    (cell.text or "").strip()
                    for cell in row.cells
                ]

                rows.append(
                    " | ".join(cells)
                )

            if rows:
                output.append(
                    f"[TABLE {table_index}]\n"
                    + "\n".join(rows)
                )

        return output

    @staticmethod
    def _docx_has_images(
        document: Any,
    ) -> bool:
        return (
            PageExtractor._docx_image_count(
                document
            )
            > 0
        )

    @staticmethod
    def _docx_image_count(
        document: Any,
    ) -> int:
        try:
            return len(
                document.inline_shapes
            )
        except Exception:
            return 0

    # ================================================================
    # Result helpers
    # ================================================================

    def _build_result(
        self,
        pages: list[PageResult],
        *,
        total_pages: int,
        failed_pages: int,
        document_type: str,
        filename: Optional[str],
        warnings: list[str],
    ) -> PageExtractionResult:
        pages = sorted(
            pages,
            key=lambda page: page.page_number,
        )

        has_text = any(
            page.has_text
            for page in pages
        )

        has_images = any(
            page.has_images
            for page in pages
        )

        # A document is considered likely scanned when it contains
        # images but little/no native text.
        text_characters = sum(
            page.character_count
            for page in pages
        )

        is_scanned = (
            has_images
            and text_characters
            < self.config.min_text_length
        )

        if not pages:
            status = PageExtractionStatus.NO_PAGES
        elif failed_pages == total_pages:
            status = PageExtractionStatus.FAILED
        elif failed_pages > 0:
            status = PageExtractionStatus.PARTIAL
        else:
            status = PageExtractionStatus.SUCCESS

        return PageExtractionResult(
            status=status,
            pages=pages,
            page_count=total_pages,
            extracted_pages=len(pages),
            failed_pages=failed_pages,
            document_type=document_type,
            filename=filename,
            has_text=has_text,
            has_images=has_images,
            is_scanned=is_scanned,
            warnings=warnings,
            metadata={
                "text_characters": text_characters,
                "text_words": sum(
                    page.word_count
                    for page in pages
                ),
            },
        )

    @staticmethod
    def _disabled_result(
        document_type: str,
        filename: Optional[str],
    ) -> PageExtractionResult:
        return PageExtractionResult(
            status=PageExtractionStatus.NO_PAGES,
            document_type=document_type,
            filename=filename,
            warnings=[
                f"{document_type.upper()} page extraction is disabled."
            ],
        )

    # ================================================================
    # Input / normalization
    # ================================================================

    def _load_content(
        self,
        source: str | Path | bytes | bytearray,
    ) -> bytes:
        if isinstance(
            source,
            (bytes, bytearray),
        ):
            return bytes(source)

        path = Path(source)

        if not path.exists():
            raise PageExtractionError(
                f"File does not exist: {path}",
                code="FILE_NOT_FOUND",
            )

        if not path.is_file():
            raise PageExtractionError(
                f"Source is not a file: {path}",
                code="INVALID_FILE",
            )

        try:
            return path.read_bytes()
        except OSError as exc:
            raise PageExtractionError(
                f"Unable to read file: {path}",
                code="FILE_READ_ERROR",
                cause=exc,
            ) from exc

    def _validate_content(
        self,
        content: bytes,
    ) -> None:
        if not content:
            raise PageExtractionError(
                "Document content is empty.",
                code="EMPTY_CONTENT",
            )

        if (
            len(content)
            > self.config.max_document_bytes
        ):
            raise PageExtractionError(
                "Document exceeds page-extraction size limit.",
                code="DOCUMENT_TOO_LARGE",
            )

    @staticmethod
    def _resolve_filename(
        source: str | Path | bytes | bytearray,
        filename: Optional[str],
    ) -> Optional[str]:
        if filename:
            return Path(filename).name

        if isinstance(
            source,
            (str, Path),
        ):
            return Path(source).name

        return None

    @staticmethod
    def _extension(
        filename: Optional[str],
    ) -> str:
        if not filename:
            return ""

        return Path(filename).suffix.lower()

    def _normalize_text(
        self,
        text: str,
    ) -> str:
        if not text:
            return ""

        text = text.replace(
            "\x00",
            "",
        )

        if self.config.normalize_whitespace:
            text = re.sub(
                r"[ \t]+",
                " ",
                text,
            )

            text = re.sub(
                r"\n[ \t]+",
                "\n",
                text,
            )

            text = re.sub(
                r"\n{3,}",
                "\n\n",
                text,
            )

        return text.strip()


def extract_pages(
    source: str | Path | bytes | bytearray,
    *,
    filename: Optional[str] = None,
    document_type: Optional[str] = None,
) -> PageExtractionResult:
    """Convenience API for the ingestion pipeline."""

    extractor = PageExtractor()

    return extractor.extract(
        source,
        filename=filename,
        document_type=document_type,
    )
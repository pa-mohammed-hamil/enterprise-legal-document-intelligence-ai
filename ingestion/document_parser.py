"""
Enterprise legal document parser.

Path:
    backend/ingestion/document_parser.py

Responsibilities:
    - Detect supported document formats.
    - Validate uploaded documents.
    - Dispatch to format-specific parsers.
    - Normalize parser output into a common structure.
    - Preserve pages, paragraphs, metadata, and source information.
    - Provide deterministic parser errors.

Non-responsibilities:
    - OCR implementation.
    - Document normalization.
    - Chunking.
    - Embedding.
    - Classification.
    - Risk analysis.
    - Legal interpretation.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path
from typing import Any, BinaryIO, Callable, Optional
import hashlib
import io
import json
import logging
import mimetypes
import re


logger = logging.getLogger(__name__)


# ======================================================================
# Enums
# ======================================================================


class DocumentFormat(str, Enum):
    PDF = "pdf"
    DOCX = "docx"
    DOC = "doc"
    TXT = "txt"
    MD = "md"
    HTML = "html"
    JSON = "json"
    UNKNOWN = "unknown"


class ParseStatus(str, Enum):
    SUCCESS = "success"
    PARTIAL = "partial"
    FAILED = "failed"


class ParserErrorCode(str, Enum):
    UNSUPPORTED_FORMAT = "unsupported_format"
    EMPTY_DOCUMENT = "empty_document"
    INVALID_DOCUMENT = "invalid_document"
    FILE_TOO_LARGE = "file_too_large"
    PARSE_FAILED = "parse_failed"
    ENCODING_ERROR = "encoding_error"
    DEPENDENCY_MISSING = "dependency_missing"


# ======================================================================
# Data structures
# ======================================================================


@dataclass(slots=True)
class DocumentPage:
    """Normalized representation of one document page."""

    page_number: int

    text: str

    metadata: dict[str, Any] = field(
        default_factory=dict
    )


@dataclass(slots=True)
class DocumentSection:
    """Logical section extracted from a document."""

    title: Optional[str]

    text: str

    page_start: Optional[int] = None
    page_end: Optional[int] = None

    level: Optional[int] = None

    metadata: dict[str, Any] = field(
        default_factory=dict
    )


@dataclass(slots=True)
class ParsedDocument:
    """Canonical parser output."""

    document_id: str

    filename: str

    format: DocumentFormat

    mime_type: Optional[str]

    status: ParseStatus

    text: str

    pages: list[DocumentPage] = field(
        default_factory=list
    )

    sections: list[DocumentSection] = field(
        default_factory=list
    )

    metadata: dict[str, Any] = field(
        default_factory=dict
    )

    warnings: list[str] = field(
        default_factory=list
    )

    errors: list[str] = field(
        default_factory=list
    )

    content_hash: Optional[str] = None

    page_count: int = 0

    character_count: int = 0

    word_count: int = 0


@dataclass(slots=True)
class ParserConfig:
    """Parser configuration."""

    max_file_size_bytes: int = 50 * 1024 * 1024

    encoding: str = "utf-8"

    detect_encoding: bool = True

    extract_metadata: bool = True

    extract_sections: bool = True

    preserve_pages: bool = True

    allow_empty_pages: bool = False

    calculate_hash: bool = True


@dataclass(slots=True)
class ParserError(Exception):
    """Structured parser exception."""

    code: ParserErrorCode

    message: str

    filename: Optional[str] = None

    cause: Optional[str] = None

    def __str__(self) -> str:
        if self.filename:
            return (
                f"[{self.code.value}] "
                f"{self.filename}: {self.message}"
            )

        return (
            f"[{self.code.value}] "
            f"{self.message}"
        )


# ======================================================================
# Parser
# ======================================================================


class DocumentParser:
    """
    Main document parsing entry point.

    Format-specific parsers can be injected so this class remains
    independent from particular PDF/DOCX libraries.

    Example:

        parser = DocumentParser(
            pdf_parser=my_pdf_parser,
            docx_parser=my_docx_parser,
        )

        result = parser.parse(
            file_path="contract.pdf"
        )
    """

    SUPPORTED_EXTENSIONS: dict[
        str,
        DocumentFormat,
    ] = {
        ".pdf": DocumentFormat.PDF,
        ".docx": DocumentFormat.DOCX,
        ".doc": DocumentFormat.DOC,
        ".txt": DocumentFormat.TXT,
        ".md": DocumentFormat.MD,
        ".html": DocumentFormat.HTML,
        ".htm": DocumentFormat.HTML,
        ".json": DocumentFormat.JSON,
    }

    MIME_TYPES: dict[
        DocumentFormat,
        str,
    ] = {
        DocumentFormat.PDF: "application/pdf",
        DocumentFormat.DOCX: (
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        DocumentFormat.DOC: "application/msword",
        DocumentFormat.TXT: "text/plain",
        DocumentFormat.MD: "text/markdown",
        DocumentFormat.HTML: "text/html",
        DocumentFormat.JSON: "application/json",
    }

    def __init__(
        self,
        *,
        config: Optional[ParserConfig] = None,
        pdf_parser: Optional[
            Callable[[bytes, str], Any]
        ] = None,
        docx_parser: Optional[
            Callable[[bytes, str], Any]
        ] = None,
        doc_parser: Optional[
            Callable[[bytes, str], Any]
        ] = None,
        ocr_parser: Optional[
            Callable[[bytes, str], Any]
        ] = None,
    ) -> None:
        self.config = (
            config
            or ParserConfig()
        )

        self.pdf_parser = pdf_parser
        self.docx_parser = docx_parser
        self.doc_parser = doc_parser
        self.ocr_parser = ocr_parser

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def parse(
        self,
        file_path: str | Path,
        *,
        document_id: Optional[str] = None,
        metadata: Optional[
            dict[str, Any]
        ] = None,
    ) -> ParsedDocument:
        """
        Parse a document from disk.
        """

        path = Path(file_path)

        if not path.exists():
            raise ParserError(
                ParserErrorCode.INVALID_DOCUMENT,
                "File does not exist.",
                str(path),
            )

        if not path.is_file():
            raise ParserError(
                ParserErrorCode.INVALID_DOCUMENT,
                "Path is not a regular file.",
                str(path),
            )

        filename = path.name

        self._validate_file_size(
            path,
            filename,
        )

        try:
            content = path.read_bytes()
        except OSError as exc:
            raise ParserError(
                ParserErrorCode.INVALID_DOCUMENT,
                "Unable to read document.",
                filename,
                str(exc),
            ) from exc

        return self.parse_bytes(
            content,
            filename=filename,
            document_id=document_id,
            metadata=metadata,
        )

    def parse_bytes(
        self,
        content: bytes,
        *,
        filename: str,
        document_id: Optional[str] = None,
        mime_type: Optional[str] = None,
        metadata: Optional[
            dict[str, Any]
        ] = None,
    ) -> ParsedDocument:
        """
        Parse document bytes.
        """

        self._validate_content(
            content,
            filename,
        )

        document_format = (
            self.detect_format(
                filename,
                mime_type,
                content,
            )
        )

        if document_format == DocumentFormat.UNKNOWN:
            raise ParserError(
                ParserErrorCode.UNSUPPORTED_FORMAT,
                "Unsupported document format.",
                filename,
            )

        document_id = (
            document_id
            or self.generate_document_id(
                content,
                filename,
            )
        )

        content_hash = (
            self.calculate_hash(content)
            if self.config.calculate_hash
            else None
        )

        try:
            result = self._dispatch(
                content,
                filename,
                document_format,
            )

            parsed = self._coerce_result(
                result,
                content=content,
                filename=filename,
                document_id=document_id,
                document_format=document_format,
                mime_type=(
                    mime_type
                    or self.MIME_TYPES.get(
                        document_format
                    )
                ),
                content_hash=content_hash,
                metadata=metadata,
            )

            self._finalize(
                parsed
            )

            return parsed

        except ParserError:
            raise

        except ImportError as exc:
            raise ParserError(
                ParserErrorCode.DEPENDENCY_MISSING,
                "Required document parsing dependency is missing.",
                filename,
                str(exc),
            ) from exc

        except UnicodeDecodeError as exc:
            raise ParserError(
                ParserErrorCode.ENCODING_ERROR,
                "Unable to decode document text.",
                filename,
                str(exc),
            ) from exc

        except Exception as exc:
            logger.exception(
                "Document parsing failed: %s",
                filename,
            )

            raise ParserError(
                ParserErrorCode.PARSE_FAILED,
                "Document parsing failed.",
                filename,
                str(exc),
            ) from exc

    def parse_stream(
        self,
        stream: BinaryIO,
        *,
        filename: str,
        document_id: Optional[str] = None,
        mime_type: Optional[str] = None,
        metadata: Optional[
            dict[str, Any]
        ] = None,
    ) -> ParsedDocument:
        """
        Parse a file-like binary stream.
        """

        if not hasattr(stream, "read"):
            raise TypeError(
                "stream must provide read()."
            )

        content = stream.read()

        if not isinstance(
            content,
            bytes,
        ):
            raise TypeError(
                "Binary stream must return bytes."
            )

        return self.parse_bytes(
            content,
            filename=filename,
            document_id=document_id,
            mime_type=mime_type,
            metadata=metadata,
        )

    def detect_format(
        self,
        filename: str,
        mime_type: Optional[str] = None,
        content: Optional[bytes] = None,
    ) -> DocumentFormat:
        """
        Detect document format using extension, MIME type,
        and basic magic-byte inspection.
        """

        suffix = (
            Path(filename)
            .suffix
            .lower()
        )

        extension_format = (
            self.SUPPORTED_EXTENSIONS.get(
                suffix,
                DocumentFormat.UNKNOWN,
            )
        )

        # Strong content signatures take precedence.
        if content:
            if content.startswith(
                b"%PDF-"
            ):
                return DocumentFormat.PDF

            if content.startswith(
                b"PK"
            ) and suffix == ".docx":
                return DocumentFormat.DOCX

            if self._looks_like_html(
                content
            ):
                return DocumentFormat.HTML

        if extension_format != (
            DocumentFormat.UNKNOWN
        ):
            return extension_format

        if mime_type:
            normalized_mime = (
                mime_type.lower()
                .split(";")[0]
                .strip()
            )

            for fmt, known_mime in (
                self.MIME_TYPES.items()
            ):
                if normalized_mime == known_mime:
                    return fmt

        return DocumentFormat.UNKNOWN

    @staticmethod
    def calculate_hash(
        content: bytes,
    ) -> str:
        """
        SHA-256 content hash.

        Used for deduplication, provenance, and auditability.
        """

        return hashlib.sha256(
            content
        ).hexdigest()

    @staticmethod
    def generate_document_id(
        content: bytes,
        filename: str,
    ) -> str:
        """
        Stable document identifier based on content and filename.
        """

        digest = hashlib.sha256()

        digest.update(content)
        digest.update(
            filename.encode(
                "utf-8",
                errors="replace",
            )
        )

        return digest.hexdigest()[:32]

    # ------------------------------------------------------------------
    # Dispatch
    # ------------------------------------------------------------------

    def _dispatch(
        self,
        content: bytes,
        filename: str,
        document_format: DocumentFormat,
    ) -> Any:
        if document_format == DocumentFormat.PDF:
            return self._parse_pdf(
                content,
                filename,
            )

        if document_format == DocumentFormat.DOCX:
            return self._parse_docx(
                content,
                filename,
            )

        if document_format == DocumentFormat.DOC:
            return self._parse_doc(
                content,
                filename,
            )

        if document_format in {
            DocumentFormat.TXT,
            DocumentFormat.MD,
        }:
            return self._parse_text(
                content,
                filename,
            )

        if document_format == DocumentFormat.HTML:
            return self._parse_html(
                content,
                filename,
            )

        if document_format == DocumentFormat.JSON:
            return self._parse_json(
                content,
                filename,
            )

        raise ParserError(
            ParserErrorCode.UNSUPPORTED_FORMAT,
            f"Unsupported format: {document_format.value}",
            filename,
        )

    # ------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------

    def _parse_pdf(
        self,
        content: bytes,
        filename: str,
    ) -> Any:
        """
        Delegate PDF parsing to pdf_processor.py.

        Expected injected parser:

            parser(content, filename)

        It may return:
            - ParsedDocument
            - dict
            - list[DocumentPage]
            - string
        """

        if self.pdf_parser is not None:
            return self.pdf_parser(
                content,
                filename,
            )

        # Optional direct fallback.
        try:
            import pypdf
        except ImportError as exc:
            raise ImportError(
                "Install pypdf or inject pdf_parser."
            ) from exc

        reader = pypdf.PdfReader(
            io.BytesIO(content)
        )

        pages: list[
            DocumentPage
        ] = []

        for index, page in enumerate(
            reader.pages,
            start=1,
        ):
            text = page.extract_text() or ""

            pages.append(
                DocumentPage(
                    page_number=index,
                    text=text,
                )
            )

        metadata = {}

        if reader.metadata:
            metadata = {
                str(key).lstrip("/"): str(value)
                for key, value in reader.metadata.items()
                if value is not None
            }

        return {
            "pages": pages,
            "metadata": metadata,
        }

    # ------------------------------------------------------------------
    # DOCX
    # ------------------------------------------------------------------

    def _parse_docx(
        self,
        content: bytes,
        filename: str,
    ) -> Any:
        if self.docx_parser is not None:
            return self.docx_parser(
                content,
                filename,
            )

        try:
            from docx import Document
        except ImportError as exc:
            raise ImportError(
                "Install python-docx or inject docx_parser."
            ) from exc

        document = Document(
            io.BytesIO(content)
        )

        paragraphs: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        tables: list[list[list[str]]] = []

        for table in document.tables:
            rows: list[list[str]] = []

            for row in table.rows:
                rows.append(
                    [
                        cell.text.strip()
                        for cell in row.cells
                    ]
                )

            tables.append(rows)

        return {
            "text": "\n\n".join(
                paragraphs
            ),
            "metadata": {
                "paragraph_count": len(
                    paragraphs
                ),
                "table_count": len(
                    tables
                ),
                "tables": tables,
            },
        }

    # ------------------------------------------------------------------
    # DOC
    # ------------------------------------------------------------------

    def _parse_doc(
        self,
        content: bytes,
        filename: str,
    ) -> Any:
        if self.doc_parser is not None:
            return self.doc_parser(
                content,
                filename,
            )

        raise ParserError(
            ParserErrorCode.DEPENDENCY_MISSING,
            "Legacy .doc parsing requires an injected doc_parser.",
            filename,
        )

    # ------------------------------------------------------------------
    # Plain text / Markdown
    # ------------------------------------------------------------------

    def _parse_text(
        self,
        content: bytes,
        filename: str,
    ) -> Any:
        encoding = self.config.encoding

        if self.config.detect_encoding:
            encoding = self._detect_text_encoding(
                content
            )

        text = content.decode(
            encoding,
            errors="strict",
        )

        return {
            "text": text,
            "metadata": {
                "encoding": encoding,
            },
        }

    # ------------------------------------------------------------------
    # HTML
    # ------------------------------------------------------------------

    def _parse_html(
        self,
        content: bytes,
        filename: str,
    ) -> Any:
        try:
            from bs4 import BeautifulSoup
        except ImportError as exc:
            raise ImportError(
                "Install beautifulsoup4 or inject an HTML parser."
            ) from exc

        soup = BeautifulSoup(
            content,
            "html.parser",
        )

        # Remove non-document content.
        for element in soup(
            [
                "script",
                "style",
                "noscript",
                "svg",
            ]
        ):
            element.decompose()

        text = soup.get_text(
            separator="\n",
            strip=True,
        )

        title = None

        if soup.title:
            title = (
                soup.title.get_text(
                    strip=True
                )
            )

        return {
            "text": text,
            "metadata": {
                "title": title,
            },
        }

    # ------------------------------------------------------------------
    # JSON
    # ------------------------------------------------------------------

    def _parse_json(
        self,
        content: bytes,
        filename: str,
    ) -> Any:
        data = json.loads(
            content.decode(
                self.config.encoding
            )
        )

        if isinstance(
            data,
            str,
        ):
            text = data

        elif isinstance(
            data,
            dict,
        ):
            # Prefer explicit document text fields.
            text = self._extract_json_text(
                data
            )

        else:
            text = json.dumps(
                data,
                ensure_ascii=False,
                indent=2,
            )

        return {
            "text": text,
            "metadata": {
                "json_type": type(data).__name__,
            },
        }

    @staticmethod
    def _extract_json_text(
        data: dict[str, Any],
    ) -> str:
        preferred_fields = (
            "text",
            "content",
            "body",
            "document_text",
            "document_content",
        )

        for field_name in preferred_fields:
            value = data.get(
                field_name
            )

            if isinstance(
                value,
                str,
            ):
                return value

        return json.dumps(
            data,
            ensure_ascii=False,
            indent=2,
        )

    # ------------------------------------------------------------------
    # Result coercion
    # ------------------------------------------------------------------

    def _coerce_result(
        self,
        result: Any,
        *,
        content: bytes,
        filename: str,
        document_id: str,
        document_format: DocumentFormat,
        mime_type: Optional[str],
        content_hash: Optional[str],
        metadata: Optional[
            dict[str, Any]
        ],
    ) -> ParsedDocument:
        if isinstance(
            result,
            ParsedDocument,
        ):
            result.document_id = document_id
            result.filename = filename
            result.format = document_format
            result.mime_type = mime_type

            if content_hash:
                result.content_hash = (
                    content_hash
                )

            result.metadata.update(
                metadata or {}
            )

            return result

        pages: list[
            DocumentPage
        ] = []

        sections: list[
            DocumentSection
        ] = []

        result_metadata: dict[
            str,
            Any,
        ] = {}

        text = ""

        if isinstance(
            result,
            str,
        ):
            text = result

        elif isinstance(
            result,
            list,
        ):
            pages = self._coerce_pages(
                result
            )

            text = self._combine_pages(
                pages
            )

        elif isinstance(
            result,
            dict,
        ):
            raw_text = result.get(
                "text",
                "",
            )

            if isinstance(
                raw_text,
                str,
            ):
                text = raw_text

            if "pages" in result:
                pages = self._coerce_pages(
                    result["pages"]
                )

                if not text:
                    text = self._combine_pages(
                        pages
                    )

            if "sections" in result:
                sections = (
                    self._coerce_sections(
                        result["sections"]
                    )
                )

            result_metadata.update(
                result.get(
                    "metadata",
                    {},
                )
            )

        else:
            raise ParserError(
                ParserErrorCode.PARSE_FAILED,
                "Parser returned an unsupported result type.",
                filename,
            )

        if not text.strip() and not pages:
            raise ParserError(
                ParserErrorCode.EMPTY_DOCUMENT,
                "No text was extracted from document.",
                filename,
            )

        return ParsedDocument(
            document_id=document_id,
            filename=filename,
            format=document_format,
            mime_type=mime_type,
            status=ParseStatus.SUCCESS,
            text=text,
            pages=pages,
            sections=sections,
            metadata={
                **result_metadata,
                **(metadata or {}),
            },
            content_hash=content_hash,
        )

    # ------------------------------------------------------------------
    # Page coercion
    # ------------------------------------------------------------------

    @staticmethod
    def _coerce_pages(
        raw_pages: Any,
    ) -> list[DocumentPage]:
        if not isinstance(
            raw_pages,
            list,
        ):
            return []

        pages: list[
            DocumentPage
        ] = []

        for index, page in enumerate(
            raw_pages,
            start=1,
        ):
            if isinstance(
                page,
                DocumentPage,
            ):
                pages.append(page)
                continue

            if isinstance(
                page,
                str,
            ):
                pages.append(
                    DocumentPage(
                        page_number=index,
                        text=page,
                    )
                )
                continue

            if isinstance(
                page,
                dict,
            ):
                pages.append(
                    DocumentPage(
                        page_number=int(
                            page.get(
                                "page_number",
                                index,
                            )
                        ),
                        text=str(
                            page.get(
                                "text",
                                "",
                            )
                        ),
                        metadata=dict(
                            page.get(
                                "metadata",
                                {},
                            )
                        ),
                    )
                )

        return pages

    @staticmethod
    def _coerce_sections(
        raw_sections: Any,
    ) -> list[DocumentSection]:
        if not isinstance(
            raw_sections,
            list,
        ):
            return []

        sections: list[
            DocumentSection
        ] = []

        for section in raw_sections:
            if isinstance(
                section,
                DocumentSection,
            ):
                sections.append(section)
                continue

            if isinstance(
                section,
                dict,
            ):
                sections.append(
                    DocumentSection(
                        title=section.get(
                            "title"
                        ),
                        text=str(
                            section.get(
                                "text",
                                "",
                            )
                        ),
                        page_start=section.get(
                            "page_start"
                        ),
                        page_end=section.get(
                            "page_end"
                        ),
                        level=section.get(
                            "level"
                        ),
                        metadata=dict(
                            section.get(
                                "metadata",
                                {},
                            )
                        ),
                    )
                )

        return sections

    @staticmethod
    def _combine_pages(
        pages: list[DocumentPage],
    ) -> str:
        return "\n\n".join(
            page.text.strip()
            for page in pages
            if page.text.strip()
        )

    # ------------------------------------------------------------------
    # Finalization
    # ------------------------------------------------------------------

    def _finalize(
        self,
        document: ParsedDocument,
    ) -> None:
        document.page_count = len(
            document.pages
        )

        document.character_count = len(
            document.text
        )

        document.word_count = len(
            re.findall(
                r"\S+",
                document.text,
            )
        )

        if (
            document.page_count == 0
            and document.text
        ):
            document.page_count = 1

        if not document.text.strip():
            document.status = (
                ParseStatus.FAILED
            )

            document.errors.append(
                "Document contains no extracted text."
            )

        elif (
            document.page_count > 0
            and any(
                not page.text.strip()
                for page in document.pages
            )
            and not self.config.allow_empty_pages
        ):
            document.status = (
                ParseStatus.PARTIAL
            )

            document.warnings.append(
                "One or more pages contain no extracted text."
            )

    # ------------------------------------------------------------------
    # Validation
    # ------------------------------------------------------------------

    def _validate_file_size(
        self,
        path: Path,
        filename: str,
    ) -> None:
        try:
            size = path.stat().st_size
        except OSError as exc:
            raise ParserError(
                ParserErrorCode.INVALID_DOCUMENT,
                "Unable to determine file size.",
                filename,
                str(exc),
            ) from exc

        if (
            size
            > self.config.max_file_size_bytes
        ):
            raise ParserError(
                ParserErrorCode.FILE_TOO_LARGE,
                (
                    "Document exceeds maximum "
                    f"size of "
                    f"{self.config.max_file_size_bytes} bytes."
                ),
                filename,
            )

    def _validate_content(
        self,
        content: bytes,
        filename: str,
    ) -> None:
        if not isinstance(
            content,
            bytes,
        ):
            raise TypeError(
                "content must be bytes"
            )

        if not content:
            raise ParserError(
                ParserErrorCode.EMPTY_DOCUMENT,
                "Document is empty.",
                filename,
            )

        if (
            len(content)
            > self.config.max_file_size_bytes
        ):
            raise ParserError(
                ParserErrorCode.FILE_TOO_LARGE,
                (
                    "Document exceeds maximum "
                    f"size of "
                    f"{self.config.max_file_size_bytes} bytes."
                ),
                filename,
            )

    # ------------------------------------------------------------------
    # Encoding
    # ------------------------------------------------------------------

    def _detect_text_encoding(
        self,
        content: bytes,
    ) -> str:
        if content.startswith(
            b"\xef\xbb\xbf"
        ):
            return "utf-8-sig"

        if content.startswith(
            b"\xff\xfe"
        ):
            return "utf-16"

        if content.startswith(
            b"\xfe\xff"
        ):
            return "utf-16"

        # UTF-8 is preferred.
        try:
            content.decode(
                "utf-8"
            )
            return "utf-8"
        except UnicodeDecodeError:
            pass

        # Conservative fallback.
        return "latin-1"

    # ------------------------------------------------------------------
    # HTML detection
    # ------------------------------------------------------------------

    @staticmethod
    def _looks_like_html(
        content: bytes,
    ) -> bool:
        sample = content[
            :4096
        ].lower()

        return (
            b"<html" in sample
            or b"<!doctype html" in sample
            or b"<body" in sample
        )


# ======================================================================
# Convenience functions
# ======================================================================


_default_parser = DocumentParser()


def parse_document(
    file_path: str | Path,
    *,
    document_id: Optional[str] = None,
    metadata: Optional[
        dict[str, Any]
    ] = None,
) -> ParsedDocument:
    """
    Convenience function for standard parsing.
    """

    return _default_parser.parse(
        file_path,
        document_id=document_id,
        metadata=metadata,
    )


def parse_document_bytes(
    content: bytes,
    *,
    filename: str,
    document_id: Optional[str] = None,
    mime_type: Optional[str] = None,
    metadata: Optional[
        dict[str, Any]
    ] = None,
) -> ParsedDocument:
    """
    Convenience function for parsing bytes.
    """

    return _default_parser.parse_bytes(
        content,
        filename=filename,
        document_id=document_id,
        mime_type=mime_type,
        metadata=metadata,
    )


__all__ = [
    "DocumentFormat",
    "ParseStatus",
    "ParserErrorCode",
    "DocumentPage",
    "DocumentSection",
    "ParsedDocument",
    "ParserConfig",
    "ParserError",
    "DocumentParser",
    "parse_document",
    "parse_document_bytes",
]